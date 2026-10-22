"""Stage A 大纲抽取：优先文档自身结构（PDF 书签/标题、PPTX 标题占位符），
不足时一次轻量 LLM fallback，再跨文档按规范化标题合并。

设计约束（评审 §4.6 + GPT 方案 v1.0 修正案 C）：
- 容器是结构不是知识点，不要求来源；文档文件名不作为容器标题；
- 两级大纲（chapter/section），页码区间换算自文档自身顺序；
- 结构不足且 LLM 也给不出有效章节 → 显式失败（KNOWLEDGE_OUTLINE_FAILED），
  不造假结构。
"""

from __future__ import annotations

import hashlib
import logging
import math
import re
import statistics
import unicodedata
from collections import Counter
from dataclasses import dataclass, field, replace
from typing import Any, Protocol

import fitz
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

from app.core.config import get_settings
from app.core.errors import AppError
from app.llm.schemas import KnowledgeOutlineV5
from app.llm.prompts.knowledge_outline_v5 import (
    OUTLINE_PROMPT_VERSION,
    SYSTEM_PROMPT,
)
from app.storage.local import LocalStorage

logger = logging.getLogger(__name__)

# 大纲解析逻辑的结构版本；改变候选/聚合规则时必须递增，section 缓存键包含它，
# 避免"章节边界变了但文本没变"时误命中旧抽取结果。
# v3（2026-10-22）：新增文档级伞形章 wrap，层级从"文档内小节当章"改为"章→节"。
OUTLINE_PARSER_VERSION = "outline-parser-v3"

MAX_TOP_LEVEL_CHAPTERS = 20
PAGES_PER_TOP_LEVEL_CHAPTER = 12
MIN_PAGES_FOR_DENSITY_CHECK = 20
# 一级节点数 / 页数 超过该比例 → 判为"每页一个标题"的可疑大纲
MAX_TOP_LEVEL_DENSITY = 0.5
MAX_INVALID_TITLE_SHARE = 0.1
MAX_LEVEL_ONE_WITHOUT_SECTIONS = 8
HEADING_SIZE_DELTA = 1.5
MAX_CANDIDATES_PER_PAGE = 3
MAX_CANDIDATE_LINES_PER_PAGE = 8
RUNNING_HEADER_SHARE = 0.3
RUNNING_HEADER_MIN_PAGES = 4
OUTLINE_FALLBACK_ATTEMPTS = 2

REJECTED_OUTLINE_MARKERS = (
    "table of contents",
    "contents",
    "agenda",
    "acknowledgement",
    "acknowledgment",
    "copyright",
    "course outline",
    "chapter outline",
    "目录",
    "致谢",
)

_CHAPTER_RE = re.compile(
    r"^\s*(第[0-9一二三四五六七八九十百千]+章|chapter\s+\d+)",
    re.IGNORECASE,
)
_NUMBERED_RE = re.compile(r"^\s*(\d{1,2}(?:\.\d{1,2}){0,3})[\s.、:：\-–—]")
_TOC_TAIL_RE = re.compile(r"[\s.·…_\-–—]{2,}\d{1,4}$")
# 纯页码/纯编号残渣：如 "1"、"4.2"、"第 3 章"、"Chapter 4"，不能作为容器标题
_NUMERIC_ONLY_RE = re.compile(r"^[\s0-9ivxIVX.．、:：\-–—_/\\()（）]{1,8}$")
_NUMBERING_ONLY_RE = re.compile(
    r"^(第\s*\d+\s*[章节]|chapter\s*\d+|\d{1,2}(?:\.\d{1,2}){0,3})$", re.IGNORECASE
)
_PUNCT_ONLY_RE = re.compile(r"^[\s\W_]+$")


@dataclass(frozen=True)
class OutlineSection:
    """单文档内的章节条目；page 区间为闭区间。"""

    document_id: str
    title: str
    level: int  # 1=chapter, 2=section
    page_start: int
    page_end: int


@dataclass
class MergedSection:
    """跨文档合并后的章节节点；page_ranges 记录每份文档各自的页码区间。"""

    title: str
    level: int
    page_ranges: list[tuple[str, int, int]] = field(default_factory=list)
    children: list[MergedSection] = field(default_factory=list)
    # 运行期索引：规范化标题 → 子节点，用于跨文档同名小节合并
    children_index: dict[str, MergedSection] = field(default_factory=dict)


@dataclass
class OutlineCallRecord:
    purpose: str = "outline_fallback"
    document_id: str | None = None
    model_name: str = ""
    prompt_version: str = OUTLINE_PROMPT_VERSION
    latency_ms: int | None = None
    input_chars: int | None = None
    output_chars: int | None = None
    attempt: int | None = None
    status: str = "succeeded"
    error_code: str | None = None


@dataclass(frozen=True)
class HeadingCandidate:
    """单页的标题候选；只描述"这一行像标题"，不代表它是章节。"""

    page_number: int
    title: str
    level: int
    source: str  # "numbered" | "font"
    size: float = 0.0


@dataclass(frozen=True)
class OutlineQuality:
    """结构质量判定结果；三条解析路径（书签/字号启发/LLM）共用同一判定。"""

    ok: bool
    reasons: tuple[str, ...] = ()
    top_level: int = 0
    level_two: int = 0
    limit: int = MAX_TOP_LEVEL_CHAPTERS
    sections: int = 0
    invalid_titles: int = 0
    page_count: int = 0

    def describe(self) -> str:
        return "ok" if self.ok else ", ".join(self.reasons)

    def feedback(self) -> str:
        """给 LLM 修复尝试的结构性反馈（不含原文，仅数量与形态）。"""
        parts: list[str] = []
        if self.top_level > self.limit:
            parts.append(
                f"the previous attempt returned {self.top_level} level-1 chapters, "
                f"but at most {self.limit} are allowed"
            )
        if "per_page_headings" in self.reasons:
            parts.append(
                "the outline looked like one chapter per page; merge consecutive "
                "pages into real chapters"
            )
        if "no_hierarchy" in self.reasons:
            parts.append(
                "no level-2 sections were returned; group pages into chapters and "
                "put sections inside them"
            )
        if "invalid_titles" in self.reasons:
            parts.append(
                f"{self.invalid_titles} titles were page numbers or numbering "
                "residue; use real headings from the page inventory"
            )
        if "empty" in self.reasons:
            parts.append("no chapters were returned for a document that has content")
        return "; ".join(parts) or "the outline was rejected by the structure check"


def classify_heading(text: str) -> tuple[int, str] | None:
    """按编号/章节模式识别标题行，返回 (level, label)；无法识别返回 None。"""
    stripped = text.strip()
    if not stripped or len(stripped) > 120:
        return None
    if _CHAPTER_RE.match(stripped):
        match = re.match(_CHAPTER_RE, stripped)
        return 1, match.group(1) if match else stripped
    numbered = _NUMBERED_RE.match(stripped)
    if numbered:
        label = numbered.group(1)
        level = 1 if label.count(".") == 0 else 2
        return level, label
    return None


def normalize_outline_title(title: str) -> str:
    normalized = unicodedata.normalize("NFKC", title).casefold()
    normalized = " ".join(normalized.split())
    return normalized.strip(" .·:：-—_")


def clean_outline_title(raw_title: str) -> str:
    title = " ".join(raw_title.replace("\r", " ").split())
    # 去掉目录条目常见的页码引导符（如 "拥塞控制 ...... 42"）
    title = _TOC_TAIL_RE.sub("", title).strip(" .·_")
    return title[:120]


def is_rejected_outline_title(title: str) -> bool:
    lowered = title.lower()
    return any(marker in lowered for marker in REJECTED_OUTLINE_MARKERS)


def is_meaningful_outline_title(title: str) -> bool:
    """标题是否有信息量：拒绝纯页码、纯编号残渣与纯标点。

    保守规则：不按字符数一刀切（`TCP`、`概述` 都可能有效）。
    """
    text = " ".join((title or "").split())
    if not text or len(text) > 120:
        return False
    if not any(char.isalnum() for char in text):
        return False
    if _NUMERIC_ONLY_RE.match(text) or _PUNCT_ONLY_RE.match(text):
        return False
    if _NUMBERING_ONLY_RE.match(text):
        return False
    return True


def outline_top_level_limit(page_count: int) -> int:
    pages = max(0, int(page_count or 0))
    return max(
        MAX_TOP_LEVEL_CHAPTERS, math.ceil(pages / PAGES_PER_TOP_LEVEL_CHAPTER)
    )


def validate_outline_quality(
    sections: list[OutlineSection], page_count: int
) -> OutlineQuality:
    """结构质量闸门：书签、字号启发、LLM 输出都必须通过同一判定。

    拒绝"一页一个章节"的伪大纲：一级数量超上限、密度接近每页一条、
    完全没有小节层、或大量页码式标题。
    """
    top_level = sum(1 for section in sections if section.level == 1)
    level_two = sum(1 for section in sections if section.level == 2)
    invalid_titles = sum(
        1 for section in sections if not is_meaningful_outline_title(section.title)
    )
    limit = outline_top_level_limit(page_count)
    reasons: list[str] = []
    if not sections:
        reasons.append("empty")
    if top_level > limit:
        reasons.append("too_many_top_level")
    if (
        page_count >= MIN_PAGES_FOR_DENSITY_CHECK
        and top_level / max(page_count, 1) >= MAX_TOP_LEVEL_DENSITY
    ):
        reasons.append("per_page_headings")
    if level_two == 0 and top_level > MAX_LEVEL_ONE_WITHOUT_SECTIONS:
        reasons.append("no_hierarchy")
    if sections and invalid_titles / len(sections) > MAX_INVALID_TITLE_SHARE:
        reasons.append("invalid_titles")
    return OutlineQuality(
        ok=not reasons,
        reasons=tuple(reasons),
        top_level=top_level,
        level_two=level_two,
        limit=limit,
        sections=len(sections),
        invalid_titles=invalid_titles,
        page_count=max(0, int(page_count or 0)),
    )


def _fill_section_ranges(
    sections: list[OutlineSection], total_pages: int
) -> list[OutlineSection]:
    """为顺序章节条目计算 page_end：到同级或更浅层的下一条目为止。"""
    filled: list[OutlineSection] = []
    for index, section in enumerate(sections):
        boundary = total_pages
        for later in sections[index + 1 :]:
            if later.level <= section.level:
                boundary = max(section.page_start, later.page_start - 1)
                break
        filled.append(
            OutlineSection(
                document_id=section.document_id,
                title=section.title,
                level=section.level,
                page_start=section.page_start,
                page_end=max(section.page_start, boundary),
            )
        )
    return filled


def merge_document_outlines(
    doc_sections: list[tuple[Any, list[OutlineSection]]],
) -> list[MergedSection]:
    """按规范化标题路径合并多份文档的大纲；文档按传入顺序（created_at asc）。"""
    chapters: list[MergedSection] = []
    chapter_index: dict[str, MergedSection] = {}

    for document, sections in doc_sections:
        current_chapter: MergedSection | None = None
        for section in sorted(sections, key=lambda item: (item.page_start, item.level)):
            key = normalize_outline_title(section.title)
            if not key:
                continue
            if section.level == 1:
                chapter = chapter_index.get(key)
                if chapter is None:
                    chapter = MergedSection(title=section.title, level=1)
                    chapters.append(chapter)
                    chapter_index[key] = chapter
                chapter.page_ranges.append(
                    (document.id, section.page_start, section.page_end)
                )
                current_chapter = chapter
            elif current_chapter is not None:
                child = current_chapter.children_index.get(key)
                if child is None:
                    child = MergedSection(title=section.title, level=2)
                    current_chapter.children.append(child)
                    current_chapter.children_index[key] = child
                child.page_ranges.append(
                    (document.id, section.page_start, section.page_end)
                )
    return chapters


def absorb_chapter_lead_in(
    sections: list[OutlineSection], page_count: int
) -> list[OutlineSection]:
    """把每章的章首/章尾过渡页并入相邻小节，补上抽取窗口的覆盖缺口。

    大纲常见形态是"章给区间、节从标题页起"：章起始页到首个节之间（封面后的
    第一章导语、章分隔页、章概述）、节与节之间的空洞、末节之后到章结束之间的
    页面不在任何节的区间里，既进不了抽取窗口，也过不了 `_collect_leaves` 的
    来源页码校验——内容静默丢失。这里只调整页码范围：章区间补齐为
    [章起始, 下章起始-1]，章内小节按声明起点铺满整章（首节前移到章起始、中间
    节延伸到下一节起点之前、末节后延到章末尾），标题、层级、条目数量都不动。

    对书签名录这类本来就严丝合缝的大纲是恒等变换；对只有章、没有节的章节
    （整章一个窗口）也不做改动。纯函数，按 (page_start, level) 顺序输入。
    """
    total = max(0, int(page_count or 0))
    if not sections:
        return []
    if total <= 0:
        total = max(section.page_end for section in sections)

    ordered = sorted(sections, key=lambda item: (item.page_start, item.level))
    chapters = [section for section in ordered if section.level == 1]
    if not chapters:
        return list(sections)

    bounds: list[tuple[int, int]] = []
    for index, chapter in enumerate(chapters):
        start = max(1, min(chapter.page_start, total))
        next_start = (
            min(chapters[index + 1].page_start, total)
            if index + 1 < len(chapters)
            else total + 1
        )
        end = max(start, min(next_start - 1, total))
        bounds.append((start, end))

    adjusted: list[OutlineSection] = []
    children_of: dict[int, list[int]] = {}
    chapter_no = -1
    for section in ordered:
        if section.level == 1:
            chapter_no += 1
            start, end = bounds[chapter_no]
            adjusted.append(replace(section, page_start=start, page_end=end))
            children_of[chapter_no] = []
        elif section.level == 2 and chapter_no >= 0:
            adjusted.append(section)
            children_of[chapter_no].append(len(adjusted) - 1)
        else:
            adjusted.append(section)

    for chapter_no, indices in children_of.items():
        if not indices:
            continue
        start, end = bounds[chapter_no]
        for position, index in enumerate(indices):
            child = adjusted[index]
            # 章内小节按声明起点铺满整章：首节前移到章起始，中间节延伸到下一节
            # 起点之前，末节后延到章末尾。LLM 给的区间常见"节与节之间留洞"，
            # 只补章首章尾不够，这些洞同样进不了窗口、过不了来源校验。
            child_start = start if position == 0 else max(child.page_start, start)
            if position + 1 < len(indices):
                next_start = max(adjusted[indices[position + 1]].page_start, start)
                child_end = max(child_start, min(next_start - 1, end))
            else:
                child_end = end
            adjusted[index] = replace(
                child, page_start=child_start, page_end=max(child_start, child_end)
            )
    return adjusted


_UMBRELLA_FILENAME_RE = re.compile(
    r"ch\s*(\d{1,2})\s*[_\-]\s*([^_()（）]+)", re.IGNORECASE
)
_CN_DIGITS = "零一二三四五六七八九"


def _chinese_number(value: int) -> str:
    if value <= 0:
        return str(value)
    if value < 10:
        return _CN_DIGITS[value]
    if value < 20:
        tens = "十"
        rest = value - 10
        return tens + (_CN_DIGITS[rest] if rest else "")
    tens, ones = divmod(value, 10)
    return _CN_DIGITS[tens] + "十" + (_CN_DIGITS[ones] if ones else "")


def umbrella_title_from_filename(filename: str) -> str | None:
    """从 `ch3_数据链路层_v2.0(3).pdf` 这类文件名取"第N章 名称"。"""
    stem = re.sub(r"\.(pdf|pptx|ppt)$", "", filename or "", flags=re.IGNORECASE)
    match = _UMBRELLA_FILENAME_RE.search(stem)
    if not match:
        return None
    name = " ".join(match.group(2).replace("_", " ").split()).strip(" -·.")
    if not name or not is_meaningful_outline_title(name):
        return None
    return f"第{_chinese_number(int(match.group(1)))}章 {name}"


def _range_union(
    section: OutlineSection, children: list[OutlineSection], total: int
) -> tuple[int, int]:
    """小节自己的区间与子节点区间取并集，避免降级时丢页。"""
    starts = [section.page_start, *(child.page_start for child in children)]
    ends = [section.page_end, *(child.page_end for child in children)]
    start = max(1, min(min(starts), total))
    end = max(start, min(max(ends), total))
    return start, end


def wrap_document_outline(
    document: Any,
    sections: list[OutlineSection],
    *,
    root_title: str | None = None,
    page_count: int | None = None,
) -> list[OutlineSection]:
    """为"单章教材/课件"补一层文档级伞形章（章 → 节 两级不变）。

    教材每份 PDF 往往是一整章，Stage A 却常把文档内部小节当成顶级章：ch3 的
    六个一级全是节名，ch4 的一份缓存是 7 个一级、另一份是 6 个扁平一级。这里
    在合并之前按三条规则收敛：

    1. 一级标题恰好一个匹配 `第X章` / `chapter N` → 它就是伞形章，页码扩到整份
       文档；其余一级降为它的二级；被降级一级原有的二级只把页码并进父区间后丢弃；
       伞形章自己的二级全部保留。
    2. 一个都不匹配 → 用 `root_title`（LLM 路径）或文件名 `chN_名称` 造伞形章，
       原一级全部降为二级；两者都拿不到就不 wrap。
    3. 两个及以上匹配 → 本来就是多章合订本，原样返回。

    降级节点的页码取"自身 ∪ 子节点"的并集，防 LLM 页码不精确丢内容。首尾是否
    相接交给 `absorb_chapter_lead_in`（它按章把区间铺满），这里只保证章边界。
    """
    total = int(page_count or 0)
    if total <= 0:
        total = max((section.page_end for section in sections), default=0)
    if not sections or total <= 0:
        return list(sections)

    ordered = sorted(sections, key=lambda item: (item.page_start, item.level))
    chapters = [section for section in ordered if section.level == 1]
    if not chapters:
        return list(sections)
    matched = [chapter for chapter in chapters if _CHAPTER_RE.match(chapter.title)]
    if len(matched) >= 2:
        return list(sections)

    children_of: dict[int, list[OutlineSection]] = {index: [] for index in range(len(chapters))}
    chapter_no = -1
    for section in ordered:
        if section.level == 1:
            chapter_no += 1
        elif section.level == 2 and chapter_no >= 0:
            children_of[chapter_no].append(section)

    def demote(index: int) -> OutlineSection:
        chapter = chapters[index]
        start, end = _range_union(chapter, children_of[index], total)
        return OutlineSection(
            document_id=chapter.document_id,
            title=chapter.title,
            level=2,
            page_start=start,
            page_end=end,
        )

    document_id = chapters[0].document_id
    if len(matched) == 1:
        umbrella_index = chapters.index(matched[0])
        umbrella = chapters[umbrella_index]
        result: list[OutlineSection] = [
            OutlineSection(document_id, umbrella.title, 1, 1, total)
        ]
        result.extend(
            demote(index) for index in range(len(chapters)) if index != umbrella_index
        )
        result.extend(children_of[umbrella_index])
    else:
        title = ""
        if root_title:
            candidate = clean_outline_title(root_title)
            if candidate and is_meaningful_outline_title(candidate):
                title = candidate
        if not title:
            title = umbrella_title_from_filename(getattr(document, "filename", "") or "") or ""
        if not title:
            logger.info(
                "no usable umbrella title for single-chapter document",
                extra={"document_id": document_id},
            )
            return list(sections)
        result = [OutlineSection(document_id, title, 1, 1, total)]
        result.extend(demote(index) for index in range(len(chapters)))

    result.sort(key=lambda item: (item.page_start, item.level))
    broken = [
        section
        for section in result
        if section.page_start < 1
        or section.page_start > section.page_end
        or section.page_end > total
    ]
    level_two = sum(1 for section in result if section.level == 2)
    titles = {normalize_outline_title(section.title) for section in result}
    if broken or level_two == 0 or len(titles) != len(result):
        logger.warning(
            "document outline wrap rejected, keeping original sections",
            extra={
                "document_id": document_id,
                "broken_ranges": len(broken),
                "level_two": level_two,
                "duplicate_titles": len(result) - len(titles),
                "page_count": total,
            },
        )
        return list(sections)
    logger.info(
        "document outline wrapped into umbrella chapter",
        extra={
            "document_id": document_id,
            "umbrella": next(section.title for section in result if section.level == 1),
            "sections": level_two,
            "branch": "existing" if len(matched) == 1 else "created",
        },
    )
    return result


class OutlineCache(Protocol):
    """Stage A 大纲缓存接口。

    DB 实现在 knowledge_pipeline（独立短事务）；测试可用字典桩。
    """

    def get(self, cache_key: str) -> dict | None: ...

    def put(
        self,
        cache_key: str,
        *,
        document_id: str,
        content_hash: str,
        prompt_version: str,
        model_name: str,
        payload: dict,
    ) -> None: ...


def document_content_hash(document: Any) -> str:
    """文档内容指纹：页号 + 提取状态 + 原文，任一变化都会换缓存键。"""
    digest = hashlib.sha256()
    for page in sorted(
        getattr(document, "pages", None) or [],
        key=lambda item: getattr(item, "page_number", 0),
    ):
        digest.update(
            f"{getattr(page, 'page_number', 0)}|"
            f"{getattr(page, 'extraction_status', '')}|".encode("utf-8")
        )
        digest.update((getattr(page, "extracted_text", "") or "").encode("utf-8"))
        digest.update(b"\n")
    return digest.hexdigest()


def _provider_hash() -> str:
    base_url = (get_settings().llm_base_url or "").strip()
    return hashlib.sha256(base_url.encode("utf-8")).hexdigest()[:16]


def outline_cache_key(document_id: str, content_hash: str, model_name: str) -> str:
    payload = (
        f"v5-outline|{OUTLINE_PARSER_VERSION}|{OUTLINE_PROMPT_VERSION}|"
        f"{document_id}|{content_hash}|{model_name}|{_provider_hash()}"
    )
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


class KnowledgeOutlineService:
    def __init__(
        self,
        storage: LocalStorage | None = None,
        cache: OutlineCache | None = None,
    ) -> None:
        self.storage = storage or LocalStorage()
        self.cache = cache
        # 每份 PDF 只解析一次行结构；字号候选同时服务大纲启发与 page_inventory
        self._candidate_cache: dict[str, dict[int, list[HeadingCandidate]]] = {}

    def build_course_outline(
        self,
        documents: list[Any],
        *,
        adapter: Any,
        course_name: str,
        reference_outline: str | None = None,
        call_log: list[OutlineCallRecord] | None = None,
        budget: Any | None = None,
    ) -> list[MergedSection]:
        doc_sections: list[tuple[Any, list[OutlineSection]]] = []
        for document in documents:
            sections, root_title = self._resolve_document_outline(
                adapter,
                document,
                course_name=course_name,
                reference_outline=reference_outline,
                call_log=call_log,
                budget=budget,
            )
            # 文档级伞形章：单章教材被拆成多个"一级节"时补一层章
            sections = wrap_document_outline(
                document,
                sections,
                root_title=root_title,
                page_count=self._document_page_count(document),
            )
            # 章首/章尾过渡页并入相邻小节：窗口规划与来源校验共用这一份区间
            sections = absorb_chapter_lead_in(
                sections, self._document_page_count(document)
            )
            doc_sections.append((document, sections))
        return merge_document_outlines(doc_sections)

    def _resolve_document_outline(
        self,
        adapter: Any,
        document: Any,
        *,
        course_name: str,
        reference_outline: str | None,
        call_log: list[OutlineCallRecord] | None,
        budget: Any | None,
    ) -> tuple[list[OutlineSection], str | None]:
        """确定性大纲必须过结构闸门，否则走 LLM fallback（同样过闸）。"""
        sections, quality = self._deterministic_outline(document)
        if sections and quality.ok:
            return sections, None
        logger.info(
            "outline quality gate rejected deterministic outline",
            extra={
                "document_id": document.id,
                "reasons": list(quality.reasons),
                "top_level": quality.top_level,
                "limit": quality.limit,
                "page_count": quality.page_count,
            },
        )
        return self._llm_outline_fallback(
            adapter,
            document,
            course_name=course_name,
            reference_outline=reference_outline,
            call_log=call_log,
            budget=budget,
            rejected=quality,
        )

    def _deterministic_outline(
        self, document: Any
    ) -> tuple[list[OutlineSection], OutlineQuality]:
        total_pages = self._document_page_count(document)
        key = (document.storage_key or "").lower()
        try:
            if key.endswith(".pdf"):
                return self._pdf_outline_with_gate(document, total_pages)
            if key.endswith(".pptx"):
                sections = self._pptx_outline(document)
                return sections, validate_outline_quality(sections, total_pages)
        except AppError:
            raise
        except Exception:
            # 结构解析失败不是构建失败：交给 LLM fallback 兜底
            logger.warning(
                "deterministic outline parsing failed",
                extra={"document_id": getattr(document, "id", None)},
                exc_info=True,
            )
        return [], validate_outline_quality([], total_pages)

    def _document_page_count(self, document: Any) -> int:
        pages = [
            getattr(page, "page_number", 0) or 0
            for page in (getattr(document, "pages", None) or [])
        ]
        return max(pages) if pages else 0

    # ------------------------------------------------------------------ PDF

    def _pdf_outline_with_gate(
        self, document: Any, total_pages: int
    ) -> tuple[list[OutlineSection], OutlineQuality]:
        """书签优先；书签不成章节结构时再试字号/编号启发；两者都过闸才算数。"""
        toc_sections = self._pdf_toc_outline(document, total_pages)
        toc_quality = validate_outline_quality(toc_sections, total_pages)
        if toc_sections and toc_quality.ok:
            return toc_sections, toc_quality
        heading_sections = self._pdf_heading_outline(document, total_pages)
        heading_quality = validate_outline_quality(heading_sections, total_pages)
        if heading_sections and heading_quality.ok:
            return heading_sections, heading_quality
        candidates = [
            (toc_sections, toc_quality),
            (heading_sections, heading_quality),
        ]
        candidates = [item for item in candidates if item[0]]
        if not candidates:
            return [], validate_outline_quality([], total_pages)
        # 都不过闸：返回一级更少的那份，供诊断与 fallback 反馈参考
        return min(candidates, key=lambda item: (item[1].top_level, -item[1].level_two))

    def _pdf_outline(self, document: Any) -> list[OutlineSection]:
        """兼容单测/调试入口：不做质量闸门，返回确定性候选。"""
        total_pages = self._document_page_count(document) or self._pdf_page_count(document)
        toc_sections = self._pdf_toc_outline(document, total_pages)
        if toc_sections:
            return toc_sections
        return self._pdf_heading_outline(document, total_pages)

    def _pdf_page_count(self, document: Any) -> int:
        path = self.storage.resolve_path(document.storage_key)
        with fitz.open(path) as pdf:
            return pdf.page_count

    def _pdf_toc_outline(
        self, document: Any, total_pages: int
    ) -> list[OutlineSection]:
        path = self.storage.resolve_path(document.storage_key)
        with fitz.open(path) as pdf:
            effective_pages = max(total_pages, pdf.page_count)
            toc = pdf.get_toc(simple=True) or []

        sections: list[OutlineSection] = []
        for level, raw_title, page in toc:
            title = clean_outline_title(raw_title or "")
            if level not in (1, 2) or not title:
                continue
            if is_rejected_outline_title(title):
                continue
            if not 1 <= page <= effective_pages:
                continue
            sections.append(OutlineSection(document.id, title, level, page, page))
        if not sections:
            return []
        return _fill_section_ranges(sections, effective_pages)

    def _heading_candidates(
        self, document: Any
    ) -> dict[int, list[HeadingCandidate]]:
        """逐页收集标题候选（编号行 + 明显大字号行），保留字号与来源。

        候选提取只回答"这一行像不像标题"，是否构成章节由质量闸门决定。
        """
        cached = self._candidate_cache.get(document.id)
        if cached is not None:
            return cached

        path = self.storage.resolve_path(document.storage_key)
        page_lines: list[list[tuple[str, float]]] = []
        with fitz.open(path) as pdf:
            for page_index in range(pdf.page_count):
                structured = pdf[page_index].get_text("dict")
                lines: list[tuple[str, float]] = []
                for block in structured.get("blocks", []):
                    for line in block.get("lines", []):
                        text = "".join(
                            span.get("text", "") for span in line.get("spans", [])
                        ).strip()
                        if not text:
                            continue
                        size = max(
                            (span.get("size", 0.0) for span in line.get("spans", [])),
                            default=0.0,
                        )
                        lines.append((text, size))
                page_lines.append(lines)

        all_sizes = [size for lines in page_lines for _, size in lines]
        body_size = statistics.median(all_sizes) if all_sizes else 0.0

        per_page: dict[int, list[HeadingCandidate]] = {}
        for page_number, lines in enumerate(page_lines, start=1):
            head = lines[:MAX_CANDIDATE_LINES_PER_PAGE]
            candidates: list[HeadingCandidate] = []
            for text, size in head:
                heading = classify_heading(text)
                if heading is None:
                    continue
                title = clean_outline_title(text)
                if not is_meaningful_outline_title(title):
                    continue
                if is_rejected_outline_title(title):
                    continue
                candidates.append(
                    HeadingCandidate(page_number, title, heading[0], "numbered", size)
                )
                break
            if body_size:
                best: HeadingCandidate | None = None
                for text, size in head:
                    if size < body_size + HEADING_SIZE_DELTA:
                        continue
                    title = clean_outline_title(text)
                    if not 2 <= len(title) <= 60:
                        continue
                    if not is_meaningful_outline_title(title):
                        continue
                    if is_rejected_outline_title(title):
                        continue
                    if best is None or size > best.size:
                        best = HeadingCandidate(page_number, title, 1, "font", size)
                if best is not None and all(
                    normalize_outline_title(best.title)
                    != normalize_outline_title(item.title)
                    for item in candidates
                ):
                    candidates.append(best)
            if candidates:
                per_page[page_number] = candidates[:MAX_CANDIDATES_PER_PAGE]

        self._candidate_cache[document.id] = per_page
        return per_page

    def _pdf_heading_outline(self, document: Any, total_pages: int) -> list[OutlineSection]:
        """无书签时的启发式：编号/章节模式行 + 明显大字号行。

        每页只取一个主候选（编号优先）；页眉类高频标题先剔除。产出是否
        真的是章节级结构由 validate_outline_quality 判定。
        """
        per_page = self._heading_candidates(document)
        if not per_page:
            return []

        counter = Counter(
            normalize_outline_title(candidate.title)
            for candidates in per_page.values()
            for candidate in candidates
        )
        threshold = max(
            RUNNING_HEADER_MIN_PAGES, int(RUNNING_HEADER_SHARE * len(per_page))
        )
        running_headers = {
            key for key, count in counter.items() if count >= threshold
        }

        candidates: list[tuple[int, str, int]] = []
        for page_number in sorted(per_page):
            primary = per_page[page_number][0]
            key = normalize_outline_title(primary.title)
            if key in running_headers:
                continue
            # 相邻页重复同一标题（页眉/续页）只保留首个
            if candidates and key == normalize_outline_title(candidates[-1][1]):
                continue
            candidates.append((page_number, primary.title, primary.level))

        if len(candidates) < 2:
            return []
        sections = [
            OutlineSection(document.id, title, level, page, page)
            for page, title, level in candidates
        ]
        return _fill_section_ranges(sections, total_pages)

    # ----------------------------------------------------------------- PPTX

    def _pptx_outline(self, document: Any) -> list[OutlineSection]:
        path = self.storage.resolve_path(document.storage_key)
        presentation = Presentation(str(path))
        total_slides = len(presentation.slides)

        candidates: list[tuple[int, str, int]] = []
        for slide_index, slide in enumerate(presentation.slides, start=1):
            title = self._slide_title(slide)
            if not title:
                continue
            heading = classify_heading(title)
            if heading is not None:
                candidates.append((slide_index, title, heading[0]))
            elif len(slide.shapes) <= 2:
                # 纯标题页视为章节页
                candidates.append((slide_index, title, 1))
        if len(candidates) < 2:
            return []
        sections = [
            OutlineSection(document.id, title, level, slide, slide)
            for slide, title, level in candidates
        ]
        return _fill_section_ranges(sections, total_slides)

    def _slide_title(self, slide: Any) -> str | None:
        title = None
        for shape in slide.shapes:
            if not getattr(shape, "is_placeholder", False) or not shape.has_text_frame:
                continue
            try:
                placeholder_type = shape.placeholder_format.type
            except Exception:
                continue
            if placeholder_type not in (
                PP_PLACEHOLDER.TITLE,
                PP_PLACEHOLDER.CENTER_TITLE,
            ):
                continue
            text = " ".join(shape.text_frame.text.split())
            if text:
                title = text[:120]
                break
        return title

    # ----------------------------------------------------------- LLM 兜底

    def _llm_outline_fallback(
        self,
        adapter: Any,
        document: Any,
        *,
        course_name: str,
        reference_outline: str | None,
        call_log: list[OutlineCallRecord] | None,
        budget: Any | None = None,
        rejected: OutlineQuality | None = None,
    ) -> tuple[list[OutlineSection], str | None]:
        """LLM 兜底：先查大纲缓存，输出同样必须过结构闸门。

        缓存命中时不再调用模型，保证重建/续跑的 section 路径与页码区间稳定；
        未命中且通过闸门的输出写回缓存。最多一次带反馈的修复尝试。
        返回值带上 root_title：文档级伞形章的名字优先取它。
        """
        total_pages = self._document_page_count(document)
        model_name = get_settings().llm_model or ""
        content_hash = document_content_hash(document)
        cache_key = outline_cache_key(document.id, content_hash, model_name)
        cached_outline = self._read_cached_outline(cache_key, document)
        if cached_outline is not None:
            sections = self._outline_from_llm(document, cached_outline)
            quality = validate_outline_quality(sections, total_pages)
            if sections and quality.ok:
                logger.info(
                    "outline cache hit",
                    extra={"document_id": document.id, "cache_key": cache_key},
                )
                return sections, cached_outline.root_title
            logger.info(
                "cached outline rejected by quality gate",
                extra={
                    "document_id": document.id,
                    "reasons": list(quality.reasons),
                },
            )

        feedback = (
            rejected.feedback() if rejected is not None and not rejected.ok else ""
        )
        last_quality: OutlineQuality | None = rejected
        for _attempt in range(OUTLINE_FALLBACK_ATTEMPTS):
            payload = self._outline_payload(
                document, course_name, reference_outline, feedback=feedback
            )
            timeout_seconds: int | None = None
            if budget is not None:
                budget.ensure_call_possible()
                timeout_seconds = budget.call_timeout()
            record = OutlineCallRecord(document_id=document.id)
            try:
                result = adapter.generate_structured_output(
                    prompt_version=OUTLINE_PROMPT_VERSION,
                    system_prompt=SYSTEM_PROMPT,
                    user_payload=payload,
                    response_model=KnowledgeOutlineV5,
                    timeout_seconds=timeout_seconds,
                )
            except AppError as exc:
                record.status = "failed"
                record.error_code = exc.code
                if call_log is not None:
                    call_log.append(record)
                raise
            outline: KnowledgeOutlineV5 = result.data
            record.model_name = result.summary.model_name
            record.latency_ms = result.summary.latency_ms
            record.attempt = result.summary.attempts
            record.input_chars = len(str(payload))
            record.output_chars = len(str(outline.model_dump()))
            if call_log is not None:
                call_log.append(record)

            sections = self._outline_from_llm(document, outline)
            quality = validate_outline_quality(sections, total_pages)
            if sections and quality.ok:
                self._write_cached_outline(
                    cache_key,
                    document=document,
                    content_hash=content_hash,
                    model_name=model_name,
                    outline=outline,
                )
                return sections, outline.root_title
            last_quality = quality
            feedback = quality.feedback()
            logger.info(
                "LLM outline rejected by quality gate",
                extra={
                    "document_id": document.id,
                    "reasons": list(quality.reasons),
                    "top_level": quality.top_level,
                    "limit": quality.limit,
                },
            )

        raise AppError(
            code="KNOWLEDGE_OUTLINE_FAILED",
            message=(
                "No usable chapter structure could be obtained for document "
                f"'{document.filename}'."
            ),
            status_code=502,
            details={
                "document_id": document.id,
                "reasons": list(last_quality.reasons) if last_quality else [],
                "top_level": last_quality.top_level if last_quality else 0,
                "limit": last_quality.limit if last_quality else 0,
            },
        )

    def _read_cached_outline(
        self, cache_key: str, document: Any
    ) -> KnowledgeOutlineV5 | None:
        if self.cache is None:
            return None
        try:
            payload = self.cache.get(cache_key)
        except Exception:
            logger.warning(
                "failed to read outline cache",
                extra={"document_id": document.id},
            )
            return None
        if not payload:
            return None
        try:
            return KnowledgeOutlineV5.model_validate(payload)
        except Exception:
            logger.warning(
                "cached outline payload is invalid",
                extra={"document_id": document.id},
            )
            return None

    def _write_cached_outline(
        self,
        cache_key: str,
        *,
        document: Any,
        content_hash: str,
        model_name: str,
        outline: KnowledgeOutlineV5,
    ) -> None:
        if self.cache is None:
            return
        try:
            self.cache.put(
                cache_key,
                document_id=document.id,
                content_hash=content_hash,
                prompt_version=OUTLINE_PROMPT_VERSION,
                model_name=model_name,
                payload=outline.model_dump(),
            )
        except Exception:
            logger.warning(
                "failed to persist outline cache",
                extra={"document_id": document.id},
            )

    def _outline_payload(
        self,
        document: Any,
        course_name: str,
        reference_outline: str | None,
        *,
        feedback: str = "",
    ) -> dict[str, Any]:
        inventory, truncated = self._page_inventory(document)
        total_pages = max(
            [getattr(page, "page_number", 0) or 0 for page in document.pages], default=1
        )
        limit = outline_top_level_limit(total_pages)
        rules = [
            "Use only page numbers that appear in page_inventory.",
            "Level 1 = chapter, level 2 = section with a valid parent_ref.",
            "Do not use the filename as a section title.",
            (
                f"Return at most {limit} level-1 chapters and merge thin pages "
                "into the nearest chapter."
            ),
            (
                "Titles must come from real headings; never use bare page numbers, "
                "page-number-only or numbering-only titles."
            ),
        ]
        if truncated:
            rules.append(
                "page_inventory omits part of the document; infer page ranges "
                "between the listed pages instead of inventing entries."
            )
        if feedback:
            rules.append(f"A previous outline attempt was rejected: {feedback}.")
        return {
            "course_name": course_name,
            "reference_outline": reference_outline,
            "document": {
                "document_id": document.id,
                "filename": document.filename,
                "total_pages": total_pages,
            },
            "page_inventory": inventory,
            "inventory_truncated": truncated,
            "constraints": {"max_chapters": limit, "max_depth": 2},
            "rules": rules,
        }

    def _page_inventory(
        self,
        document: Any,
        *,
        char_budget: int = 32000,
        excerpt_chars: int = 80,
    ) -> tuple[list[dict[str, Any]], bool]:
        """页面清单：优先给真实标题（含大字号候选），其次给开头摘录。

        超长文档按字符预算截取并显式标记 inventory_truncated，不静默丢页。
        """
        per_page = (
            self._heading_candidates(document)
            if (document.storage_key or "").lower().endswith(".pdf")
            else {}
        )
        counter = Counter(
            normalize_outline_title(candidate.title)
            for candidates in per_page.values()
            for candidate in candidates
        )
        threshold = max(
            RUNNING_HEADER_MIN_PAGES, int(RUNNING_HEADER_SHARE * max(len(per_page), 1))
        )
        running_headers = {key for key, count in counter.items() if count >= threshold}

        entries: list[dict[str, Any]] = []
        for page in sorted(document.pages, key=lambda item: item.page_number):
            if page.extraction_status != "extracted":
                continue
            text = " ".join((page.extracted_text or "").split())
            if not text:
                continue
            headings: list[str] = []
            for candidate in per_page.get(page.page_number, []):
                if normalize_outline_title(candidate.title) in running_headers:
                    continue
                if candidate.title not in headings:
                    headings.append(candidate.title[:60])
            if not headings:
                headings = [
                    line.strip()
                    for line in (page.extracted_text or "").splitlines()
                    if classify_heading(line)
                ][:3]
            if headings:
                entries.append(
                    {"page_number": page.page_number, "headings": headings[:3]}
                )
            else:
                entries.append(
                    {
                        "page_number": page.page_number,
                        "headings": [],
                        "opening_excerpt": text[:excerpt_chars],
                    }
                )

        selected: list[dict[str, Any]] = []
        used_chars = 0
        for entry in entries:
            cost = (
                sum(len(heading) for heading in entry["headings"])
                + len(entry.get("opening_excerpt", ""))
                + 32
            )
            if selected and used_chars + cost > char_budget:
                break
            selected.append(entry)
            used_chars += cost
        truncated = len(selected) < len(entries)
        if truncated:
            logger.warning(
                "page inventory truncated for LLM outline",
                extra={
                    "document_id": document.id,
                    "kept": len(selected),
                    "total": len(entries),
                    "chars": used_chars,
                },
            )
        return selected, truncated

    def _outline_from_llm(self, document: Any, outline: KnowledgeOutlineV5) -> list[OutlineSection]:
        total_pages = max(
            [page.page_number for page in document.pages], default=1
        )
        chapter_refs = {
            section.ref for section in outline.sections if section.level == 1
        }
        sections: list[OutlineSection] = []
        for section in outline.sections:
            title = clean_outline_title(section.title)
            if not title or is_rejected_outline_title(title):
                continue
            page_start = min(max(section.page_start, 1), total_pages)
            page_end = min(max(section.page_end, page_start), total_pages)
            if section.level == 2 and section.parent_ref not in chapter_refs:
                continue
            sections.append(
                OutlineSection(document.id, title, section.level, page_start, page_end)
            )
        sections.sort(key=lambda item: (item.page_start, item.level))
        return sections
