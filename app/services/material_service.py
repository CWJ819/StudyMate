import hashlib
import re
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path

import fitz
from pptx import Presentation
from sqlalchemy.orm import Session

from app.core.config import get_settings
from app.core.errors import AppError
from app.db.models.document import Document, DocumentChunk, DocumentPage
from app.repositories.course_repository import CourseRepository
from app.repositories.job_repository import JobRepository
from app.repositories.knowledge_repository import KnowledgeRepository
from app.repositories.material_repository import MaterialRepository
from app.schemas.material import (
    ChunkDetailResponse,
    ChunkSearchResult,
    DocumentParsePayload,
    MaterialDetailResponse,
    MaterialPageResponse,
    MaterialSummaryResponse,
    MaterialUploadResponse,
)
from app.storage.local import LocalStorage

DOCUMENT_STATUS_QUEUED = "queued"
DOCUMENT_STATUS_PROCESSING = "processing"
DOCUMENT_STATUS_COMPLETED = "completed"
DOCUMENT_STATUS_COMPLETED_WITH_WARNINGS = "completed_with_warnings"
DOCUMENT_STATUS_FAILED = "failed"
DOCUMENT_STATUS_ARCHIVED = "archived"

PAGE_STATUS_EXTRACTED = "extracted"
PAGE_STATUS_EMPTY = "empty"
PAGE_STATUS_NO_EXTRACTABLE_TEXT = "no_extractable_text"

ALLOWED_EXTENSIONS = {
    ".pdf": {"application/pdf", "application/octet-stream"},
    ".pptx": {
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        "application/octet-stream",
    },
}


@dataclass(frozen=True)
class ParsedPage:
    page_number: int
    extracted_text: str | None
    extraction_status: str
    error_code: str | None


@dataclass(frozen=True)
class ParsedChunk:
    page_start: int
    page_end: int
    text: str
    content_hash: str


@dataclass(frozen=True)
class ParseResult:
    status: str
    error_code: str | None
    pages: list[ParsedPage]
    chunks: list[ParsedChunk]


class MaterialService:
    def __init__(
        self,
        *,
        course_repository: CourseRepository | None = None,
        job_repository: JobRepository | None = None,
        knowledge_repository: KnowledgeRepository | None = None,
        material_repository: MaterialRepository | None = None,
        storage: LocalStorage | None = None,
    ) -> None:
        self.course_repository = course_repository or CourseRepository()
        self.job_repository = job_repository or JobRepository()
        self.knowledge_repository = knowledge_repository or KnowledgeRepository()
        self.material_repository = material_repository or MaterialRepository()
        self.storage = storage or LocalStorage()

    def upload_material(
        self,
        session: Session,
        *,
        course_id: str,
        filename: str,
        content_type: str | None,
        data: bytes,
    ) -> MaterialUploadResponse:
        course = self.course_repository.get_course(session, course_id)
        if course is None:
            raise AppError(
                code="RESOURCE_NOT_FOUND",
                message="Course not found.",
                status_code=404,
                details={"course_id": course_id},
            )
        suffix = Path(filename).suffix.lower()
        self._validate_upload(filename=filename, suffix=suffix, content_type=content_type, data=data)

        stored_file = self.storage.save_bytes(suffix=suffix, data=data)
        document = self.material_repository.create_document(
            session,
            course_id=course_id,
            filename=filename,
            storage_key=stored_file.storage_key,
            mime_type=content_type or self._default_mime_type(suffix),
            size_bytes=stored_file.size_bytes,
            status=DOCUMENT_STATUS_QUEUED,
        )
        course.material_revision += 1
        job = self.job_repository.create_job(
            session,
            job_type="document.parse",
            payload_json=DocumentParsePayload(document_id=document.id).model_dump(),
            max_attempts=get_settings().job_max_attempts,
        )
        session.commit()

        return MaterialUploadResponse(
            document_id=document.id,
            job_id=job.id,
            status=document.status,
            error_code=document.error_code,
        )

    def list_materials(self, session: Session, course_id: str) -> list[MaterialSummaryResponse]:
        self._require_course(session, course_id)
        documents = self.material_repository.list_documents(session, course_id)
        return [self._to_material_summary(document) for document in documents]

    def get_material(self, session: Session, document_id: str) -> MaterialDetailResponse:
        document = self.material_repository.get_document(session, document_id)
        if document is None:
            raise AppError(
                code="RESOURCE_NOT_FOUND",
                message="Material not found.",
                status_code=404,
                details={"material_id": document_id},
            )
        summary = self._to_material_summary(document)
        pages = [
            MaterialPageResponse(
                page_number=page.page_number,
                extraction_status=page.extraction_status,
                error_code=page.error_code,
                text_preview=self._preview(page.extracted_text),
                text_length=len(page.extracted_text or ""),
            )
            for page in document.pages
        ]
        return MaterialDetailResponse(**summary.model_dump(), pages=pages)

    def delete_material(self, session: Session, document_id: str) -> None:
        document = self.material_repository.get_document(session, document_id)
        if document is None:
            raise AppError(
                code="RESOURCE_NOT_FOUND",
                message="Material not found.",
                status_code=404,
                details={"material_id": document_id},
            )
        if document.archived_at is not None:
            return
        document.archived_at = datetime.now(UTC)
        document.status = DOCUMENT_STATUS_ARCHIVED
        course = self.course_repository.get_course(session, document.course_id)
        if course is not None:
            course.material_revision += 1
        self.knowledge_repository.mark_active_frameworks_stale(session, document.course_id)
        self.job_repository.cancel_document_jobs(session, document_id=document.id)
        session.commit()

    def search_materials(
        self, session: Session, *, course_id: str, query: str
    ) -> list[ChunkSearchResult]:
        self._require_course(session, course_id)
        chunks = self.material_repository.search_chunks(session, course_id=course_id, query=query)
        return [
            ChunkSearchResult(
                chunk_id=chunk.id,
                document_id=chunk.document_id,
                filename=chunk.document.filename,
                page_start=chunk.page_start,
                page_end=chunk.page_end,
                text=chunk.text,
            )
            for chunk in chunks
        ]

    def get_chunk(self, session: Session, chunk_id: str) -> ChunkDetailResponse:
        chunk = self.material_repository.get_chunk(session, chunk_id)
        if chunk is None:
            raise AppError(
                code="RESOURCE_NOT_FOUND",
                message="Chunk not found.",
                status_code=404,
                details={"chunk_id": chunk_id},
            )
        return ChunkDetailResponse(
            id=chunk.id,
            document_id=chunk.document_id,
            course_id=chunk.document.course_id,
            filename=chunk.document.filename,
            page_start=chunk.page_start,
            page_end=chunk.page_end,
            text=chunk.text,
        )

    def _require_course(self, session: Session, course_id: str) -> None:
        if self.course_repository.get_course(session, course_id) is None:
            raise AppError(
                code="RESOURCE_NOT_FOUND",
                message="Course not found.",
                status_code=404,
                details={"course_id": course_id},
            )

    def _validate_upload(
        self,
        *,
        filename: str,
        suffix: str,
        content_type: str | None,
        data: bytes,
    ) -> None:
        if not filename or Path(filename).name != filename or "/" in filename or "\\" in filename:
            raise AppError(
                code="INVALID_FILENAME",
                message="Uploaded filenames must not include path components.",
                status_code=400,
                details={"filename": filename},
            )
        if suffix not in ALLOWED_EXTENSIONS:
            raise AppError(
                code="INVALID_FILE_TYPE",
                message="Only .pdf and .pptx files are supported in the current stage.",
                status_code=400,
                details={"filename": filename},
            )
        if content_type and content_type not in ALLOWED_EXTENSIONS[suffix]:
            raise AppError(
                code="INVALID_FILE_TYPE",
                message="The uploaded file MIME type is not allowed.",
                status_code=400,
                details={"filename": filename, "mime_type": content_type},
            )

        settings = get_settings()
        if len(data) > settings.max_upload_bytes:
            raise AppError(
                code="FILE_TOO_LARGE",
                message="The uploaded file exceeds the configured size limit.",
                status_code=400,
            )

    def _parse_file(self, *, path: Path, suffix: str) -> ParseResult:
        if suffix == ".pdf":
            return self._parse_pdf(path)
        if suffix == ".pptx":
            return self._parse_pptx(path)
        raise AppError(
            code="INVALID_FILE_TYPE",
            message="Unsupported file type.",
            status_code=400,
        )

    def _parse_pdf(self, path: Path) -> ParseResult:
        settings = get_settings()
        document = fitz.open(path)
        if document.page_count > settings.max_document_pages:
            raise AppError(
                code="DOCUMENT_TOO_MANY_PAGES",
                message="The uploaded document exceeds the configured page limit.",
                status_code=400,
            )

        pages: list[ParsedPage] = []
        chunks: list[ParsedChunk] = []
        extracted_page_count = 0

        for index, page in enumerate(document, start=1):
            text = page.get_text("text").strip()
            if text:
                extracted_page_count += 1
                pages.append(
                    ParsedPage(
                        page_number=index,
                        extracted_text=text,
                        extraction_status=PAGE_STATUS_EXTRACTED,
                        error_code=None,
                    )
                )
                chunks.extend(self._chunk_page_text(index, text))
            else:
                pages.append(
                    ParsedPage(
                        page_number=index,
                        extracted_text=None,
                        extraction_status=PAGE_STATUS_NO_EXTRACTABLE_TEXT,
                        error_code="NO_EXTRACTABLE_TEXT",
                    )
                )

        if extracted_page_count == 0:
            return ParseResult(
                status=DOCUMENT_STATUS_FAILED,
                error_code="NO_EXTRACTABLE_TEXT",
                pages=pages,
                chunks=[],
            )
        if extracted_page_count < len(pages):
            return ParseResult(
                status=DOCUMENT_STATUS_COMPLETED_WITH_WARNINGS,
                error_code="PARTIAL_EXTRACTION",
                pages=pages,
                chunks=chunks,
            )
        return ParseResult(
            status=DOCUMENT_STATUS_COMPLETED,
            error_code=None,
            pages=pages,
            chunks=chunks,
        )

    def _parse_pptx(self, path: Path) -> ParseResult:
        settings = get_settings()
        presentation = Presentation(path)
        slide_count = len(presentation.slides)
        if slide_count > settings.max_document_pages:
            raise AppError(
                code="DOCUMENT_TOO_MANY_PAGES",
                message="The uploaded document exceeds the configured page limit.",
                status_code=400,
            )

        pages: list[ParsedPage] = []
        chunks: list[ParsedChunk] = []

        for index, slide in enumerate(presentation.slides, start=1):
            slide_text = self._extract_slide_text(slide)
            if slide_text:
                pages.append(
                    ParsedPage(
                        page_number=index,
                        extracted_text=slide_text,
                        extraction_status=PAGE_STATUS_EXTRACTED,
                        error_code=None,
                    )
                )
                chunks.extend(self._chunk_page_text(index, slide_text))
            else:
                pages.append(
                    ParsedPage(
                        page_number=index,
                        extracted_text=None,
                        extraction_status=PAGE_STATUS_EMPTY,
                        error_code=None,
                    )
                )

        return ParseResult(
            status=DOCUMENT_STATUS_COMPLETED,
            error_code=None,
            pages=pages,
            chunks=chunks,
        )

    def _extract_slide_text(self, slide) -> str:
        texts: list[str] = []
        for shape in slide.shapes:
            text = getattr(shape, "text", None)
            if text:
                normalized = self._normalize_text(text)
                if normalized:
                    texts.append(normalized)
        return "\n\n".join(texts)

    def _chunk_page_text(self, page_number: int, text: str) -> list[ParsedChunk]:
        paragraphs = [self._normalize_text(part) for part in re.split(r"\n\s*\n", text)]
        paragraphs = [paragraph for paragraph in paragraphs if paragraph]
        if not paragraphs:
            paragraphs = [self._normalize_text(text)]

        chunks: list[ParsedChunk] = []
        current_parts: list[str] = []
        current_length = 0
        max_chunk_length = 1200

        for paragraph in paragraphs:
            addition = len(paragraph) + (2 if current_parts else 0)
            if current_parts and current_length + addition > max_chunk_length:
                chunk_text = "\n\n".join(current_parts)
                chunks.append(
                    ParsedChunk(
                        page_start=page_number,
                        page_end=page_number,
                        text=chunk_text,
                        content_hash=hashlib.sha256(chunk_text.encode("utf-8")).hexdigest(),
                    )
                )
                current_parts = []
                current_length = 0

            current_parts.append(paragraph)
            current_length += addition

        if current_parts:
            chunk_text = "\n\n".join(current_parts)
            chunks.append(
                ParsedChunk(
                    page_start=page_number,
                    page_end=page_number,
                    text=chunk_text,
                    content_hash=hashlib.sha256(chunk_text.encode("utf-8")).hexdigest(),
                )
            )

        return chunks

    def process_document_parse_job(self, session: Session, *, document_id: str) -> None:
        document = self.material_repository.get_document(session, document_id)
        if document is None:
            raise AppError(
                code="RESOURCE_NOT_FOUND",
                message="Material not found.",
                status_code=404,
                details={"material_id": document_id},
            )

        if document.archived_at is not None or document.status == DOCUMENT_STATUS_ARCHIVED:
            return
        if document.status in {
            DOCUMENT_STATUS_COMPLETED,
            DOCUMENT_STATUS_COMPLETED_WITH_WARNINGS,
            DOCUMENT_STATUS_FAILED,
        } and document.pages:
            return

        document.status = DOCUMENT_STATUS_PROCESSING
        document.error_code = None

        path = self.storage.resolve_path(document.storage_key)
        parse_result = self._parse_file(path=path, suffix=Path(document.filename).suffix.lower())
        if document.archived_at is not None:
            document.status = DOCUMENT_STATUS_ARCHIVED
            return
        pages = [
            DocumentPage(
                document_id=document.id,
                page_number=page.page_number,
                extracted_text=page.extracted_text,
                extraction_status=page.extraction_status,
                error_code=page.error_code,
            )
            for page in parse_result.pages
        ]
        chunks = [
            DocumentChunk(
                document_id=document.id,
                page_start=chunk.page_start,
                page_end=chunk.page_end,
                text=chunk.text,
                content_hash=chunk.content_hash,
            )
            for chunk in parse_result.chunks
        ]
        self.material_repository.replace_pages_and_chunks(
            session,
            document_id=document.id,
            pages=pages,
            chunks=chunks,
        )
        document.status = parse_result.status
        document.error_code = parse_result.error_code

        if parse_result.status == DOCUMENT_STATUS_FAILED:
            raise AppError(
                code=parse_result.error_code or "DOCUMENT_PARSE_FAILED",
                message="Document parsing failed.",
                status_code=400,
            )

    def _normalize_text(self, text: str) -> str:
        lines = [line.strip() for line in text.splitlines()]
        filtered = [line for line in lines if line]
        return "\n".join(filtered).strip()

    def _to_material_summary(self, document: Document) -> MaterialSummaryResponse:
        no_text_pages = [
            page.page_number
            for page in document.pages
            if page.extraction_status == PAGE_STATUS_NO_EXTRACTABLE_TEXT
        ]
        return MaterialSummaryResponse(
            id=document.id,
            course_id=document.course_id,
            filename=document.filename,
            mime_type=document.mime_type,
            size_bytes=document.size_bytes,
            status=document.status,
            error_code=document.error_code,
            page_count=len(document.pages),
            chunk_count=len(document.chunks),
            no_text_pages=no_text_pages,
            created_at=document.created_at,
            updated_at=document.updated_at,
            archived_at=document.archived_at,
        )

    def _preview(self, text: str | None) -> str | None:
        if not text:
            return None
        return text[:200]

    def _default_mime_type(self, suffix: str) -> str:
        if suffix == ".pdf":
            return "application/pdf"
        if suffix == ".pptx":
            return "application/vnd.openxmlformats-officedocument.presentationml.presentation"
        return "application/octet-stream"
