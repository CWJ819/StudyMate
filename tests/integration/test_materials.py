import asyncio
from pathlib import Path

import fitz
import pytest
from httpx import ASGITransport, AsyncClient
from pptx import Presentation

from app.main import app


async def _wait_for_job(client: AsyncClient, job_id: str) -> dict:
    last_body: dict | None = None
    for _ in range(50):
        response = await client.get(f"/api/jobs/{job_id}")
        assert response.status_code == 200
        last_body = response.json()
        if last_body["status"] in {"succeeded", "failed", "cancelled"}:
            return last_body
        await asyncio.sleep(0.02)
    raise AssertionError(f"job {job_id} did not finish: {last_body}")


def _create_pdf(path: Path, page_texts: list[str]) -> None:
    document = fitz.open()
    for text in page_texts:
        page = document.new_page()
        if text:
            page.insert_text((72, 72), text)
    document.save(path)
    document.close()


def _create_pptx(path: Path, slide_texts: list[str]) -> None:
    presentation = Presentation()
    title_layout = presentation.slide_layouts[1]
    blank_layout = presentation.slide_layouts[6]

    for text in slide_texts:
        if text:
            slide = presentation.slides.add_slide(title_layout)
            slide.shapes.title.text = text.split("\n", 1)[0]
            slide.placeholders[1].text = text
        else:
            presentation.slides.add_slide(blank_layout)

    presentation.save(path)


async def _create_course(client: AsyncClient, name: str = "Physics") -> str:
    response = await client.post("/api/courses", json={"name": name})
    assert response.status_code == 201
    return response.json()["id"]


@pytest.mark.anyio
async def test_pdf_material_upload_partial_extraction_and_search(
    migrated_database_url: str,
    tmp_path: Path,
) -> None:
    pdf_path = tmp_path / "lesson.pdf"
    _create_pdf(pdf_path, ["matrix spaces are important", ""])

    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://testserver") as client:
        course_id = await _create_course(client)

        with pdf_path.open("rb") as handle:
            upload_response = await client.post(
                f"/api/courses/{course_id}/materials",
                files={"file": ("lesson.pdf", handle, "application/pdf")},
            )

        assert upload_response.status_code == 201
        upload_body = upload_response.json()
        assert upload_body["job_id"] is not None
        assert upload_body["status"] == "queued"
        assert upload_body["error_code"] is None
        material_id = upload_body["document_id"]
        job = await _wait_for_job(client, upload_body["job_id"])
        assert job["status"] == "succeeded"

        list_response = await client.get(f"/api/courses/{course_id}/materials")
        assert list_response.status_code == 200
        listed = list_response.json()
        assert len(listed) == 1
        assert listed[0]["no_text_pages"] == [2]
        assert listed[0]["chunk_count"] == 1

        detail_response = await client.get(f"/api/materials/{material_id}")
        assert detail_response.status_code == 200
        detail = detail_response.json()
        assert detail["page_count"] == 2
        assert detail["pages"][0]["extraction_status"] == "extracted"
        assert detail["pages"][1]["extraction_status"] == "no_extractable_text"

        search_response = await client.get(
            f"/api/courses/{course_id}/search",
            params={"q": "matrix"},
        )
        assert search_response.status_code == 200
        search_results = search_response.json()
        assert len(search_results) == 1
        chunk_id = search_results[0]["chunk_id"]

        chunk_response = await client.get(f"/api/chunks/{chunk_id}")
        assert chunk_response.status_code == 200
        assert "matrix spaces" in chunk_response.json()["text"]


@pytest.mark.anyio
async def test_pdf_without_extractable_text_fails_cleanly(
    migrated_database_url: str,
    tmp_path: Path,
) -> None:
    pdf_path = tmp_path / "blank.pdf"
    _create_pdf(pdf_path, ["", ""])

    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://testserver") as client:
        course_id = await _create_course(client, name="Calculus")
        with pdf_path.open("rb") as handle:
            upload_response = await client.post(
                f"/api/courses/{course_id}/materials",
                files={"file": ("blank.pdf", handle, "application/pdf")},
            )

        assert upload_response.status_code == 201
        body = upload_response.json()
        assert body["status"] == "queued"
        assert body["job_id"] is not None
        job = await _wait_for_job(client, body["job_id"])
        assert job["status"] == "failed"
        assert job["error_code"] == "NO_EXTRACTABLE_TEXT"

        detail_response = await client.get(f"/api/materials/{body['document_id']}")
        assert detail_response.status_code == 200
        detail = detail_response.json()
        assert detail["page_count"] == 2
        assert detail["chunk_count"] == 0
        assert detail["no_text_pages"] == [1, 2]


@pytest.mark.anyio
async def test_pptx_upload_and_delete_material(
    migrated_database_url: str,
    tmp_path: Path,
) -> None:
    pptx_path = tmp_path / "lesson.pptx"
    _create_pptx(pptx_path, ["Vectors\nMagnitude and direction", ""])

    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://testserver") as client:
        course_id = await _create_course(client, name="Mechanics")
        with pptx_path.open("rb") as handle:
            upload_response = await client.post(
                f"/api/courses/{course_id}/materials",
                files={
                    "file": (
                        "lesson.pptx",
                        handle,
                        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
                    )
                },
            )

        assert upload_response.status_code == 201
        body = upload_response.json()
        assert body["status"] == "queued"
        material_id = body["document_id"]
        job = await _wait_for_job(client, body["job_id"])
        assert job["status"] == "succeeded"

        detail_response = await client.get(f"/api/materials/{material_id}")
        assert detail_response.status_code == 200
        detail = detail_response.json()
        assert detail["page_count"] == 2
        assert detail["pages"][0]["extraction_status"] == "extracted"
        assert detail["pages"][1]["extraction_status"] == "empty"

        delete_response = await client.delete(f"/api/materials/{material_id}")
        assert delete_response.status_code == 204

        archived_response = await client.get(f"/api/materials/{material_id}")
        assert archived_response.status_code == 200
        assert archived_response.json()["status"] == "archived"


@pytest.mark.anyio
async def test_invalid_material_type_is_rejected(
    migrated_database_url: str,
    tmp_path: Path,
) -> None:
    text_path = tmp_path / "notes.txt"
    text_path.write_text("plain text", encoding="utf-8")

    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://testserver") as client:
        course_id = await _create_course(client, name="Literature")
        with text_path.open("rb") as handle:
            response = await client.post(
                f"/api/courses/{course_id}/materials",
                files={"file": ("notes.txt", handle, "text/plain")},
            )

        assert response.status_code == 400
        assert response.json()["error"]["code"] == "INVALID_FILE_TYPE"
