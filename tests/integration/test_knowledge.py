import asyncio
from pathlib import Path

import fitz
import pytest
from pptx import Presentation
from httpx import ASGITransport, AsyncClient
from sqlalchemy import select

from app.core.errors import AppError
from app.core.config import reset_settings_cache
from app.db.models.knowledge import KnowledgeNodeRelation
from app.db.models.llm_call import LLMCall
from app.db.session import session_scope
from app.jobs.runner import get_job_runner
from app.llm.schemas import (
    RELATION_TYPES,
    KnowledgeOutlineV5,
    KnowledgeOutlineSectionV5,
    KnowledgeSectionExtractionV5,
    KnowledgeSourceCandidate,
    LLMRunSummary,
    LLMStructuredResult,
)
from app.main import app
from app.repositories.job_repository import JobRepository
from app.repositories.knowledge_repository import KnowledgeRepository


def _create_pdf(path: Path, page_texts: list[str]) -> None:
    document = fitz.open()
    for text in page_texts:
        page = document.new_page()
        if text:
            # 中文必须用内建 china-s 字体，否则提取出的文本会退化为占位符
            page.insert_text((72, 72), text, fontname="china-s")
    document.save(path)
    document.close()


def _create_pdf_with_toc(path: Path, page_texts: list[str], toc: list[list]) -> None:
    document = fitz.open()
    for text in page_texts:
        page = document.new_page()
        page.insert_text((72, 72), text, fontname="china-s")
        page.insert_text((72, 110), "本页补充正文内容，用于生成足够长度的文本块。", fontname="china-s")
    document.set_toc(toc)
    document.save(path)
    document.close()


def _create_pptx(path: Path, titles: list[str]) -> None:
    presentation = Presentation()
    title_layout = presentation.slide_layouts[0]
    content_layout = presentation.slide_layouts[1]
    for index, title in enumerate(titles, start=1):
        slide = presentation.slides.add_slide(
            title_layout if index == 1 else content_layout
        )
        slide.shapes.title.text = title
        if index > 1:
            slide.placeholders[1].text = f"supporting content for {title} with enough text."
    presentation.save(path)


def _collect_leaves(nodes: list[dict]) -> list[dict]:
    leaves: list[dict] = []
    for node in nodes:
        if node["is_leaf"]:
            leaves.append(node)
        else:
            leaves.extend(_collect_leaves(node.get("children", [])))
    return leaves


async def _wait_for_job(client: AsyncClient, job_id: str) -> dict:
    last_body: dict | None = None
    for _ in range(200):
        response = await client.get(f"/api/jobs/{job_id}")
        assert response.status_code == 200
        last_body = response.json()
        if last_body["status"] in {"succeeded", "failed", "cancelled"}:
            return last_body
        await asyncio.sleep(0.02)
    raise AssertionError(f"job {job_id} did not finish: {last_body}")


async def _create_course(client: AsyncClient, name: str = "Biology") -> str:
    response = await client.post("/api/courses", json={"name": name})
    assert response.status_code == 201
    return response.json()["id"]


async def _upload_material(
    client: AsyncClient, course_id: str, path: Path, mime: str
) -> None:
    with path.open("rb") as handle:
        response = await client.post(
            f"/api/courses/{course_id}/materials",
            files={"file": (path.name, handle, mime)},
        )
    assert response.status_code == 201
    job = await _wait_for_job(client, response.json()["job_id"])
    assert job["status"] == "succeeded"


async def _upload_pdf_material(client: AsyncClient, course_id: str, path: Path) -> None:
    await _upload_material(client, course_id, path, "application/pdf")


PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


@pytest.mark.anyio
async def test_knowledge_build_requires_llm_configuration(
    migrated_database_url: str,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setenv("LLM_BASE_URL", "")
    monkeypatch.setenv("LLM_MODEL", "")
    monkeypatch.setenv("LLM_API_KEY", "")
    reset_settings_cache()

    pdf_path = tmp_path / "lesson.pdf"
    _create_pdf(pdf_path, ["cells divide and specialize"])

    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://testserver") as client:
        course_id = await _create_course(client)
        await _upload_pdf_material(client, course_id, pdf_path)

        response = await client.post(
            f"/api/courses/{course_id}/knowledge/build",
            json={},
        )

    assert response.status_code == 400
    assert response.json()["error"]["code"] == "LLM_NOT_CONFIGURED"
    reset_settings_cache()


@pytest.mark.anyio
async def test_knowledge_build_produces_chapter_tree(
    migrated_database_url: str,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setenv("LLM_BASE_URL", "mock://local")
    monkeypatch.setenv("LLM_MODEL", "mock-stage8")
    reset_settings_cache()

    pdf_path = tmp_path / "lesson.pdf"
    _create_pdf_with_toc(
        pdf_path,
        [
            "能量与细胞的引入内容",
            "光合作用把光能转化为化学能。",
            "细胞呼吸释放储存的能量。",
            "生态系统中的能量流动。",
        ],
        [
            [1, "第1章 能量与细胞", 1],
            [2, "1.1 光合作用", 2],
            [2, "1.2 细胞呼吸", 3],
            [1, "第2章 生态系统", 4],
        ],
    )

    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://testserver") as client:
        course_id = await _create_course(client, name="Biology Review")
        await _upload_pdf_material(client, course_id, pdf_path)

        build_response = await client.post(
            f"/api/courses/{course_id}/knowledge/build",
            json={"reference_outline": "能量与细胞"},
        )
        assert build_response.status_code == 202
        build_body = build_response.json()

        build_job = await _wait_for_job(client, build_body["job_id"])
        assert build_job["status"] == "succeeded"
        # v5 进度契约：阶段字段存在、进度到 100
        assert build_job["progress"] == 100
        assert build_job["phase"] == "activate"

        versions_response = await client.get(f"/api/courses/{course_id}/knowledge/versions")
        assert versions_response.status_code == 200
        versions = versions_response.json()
        assert len(versions) == 1
        assert versions[0]["status"] == "active"

        framework_response = await client.get(f"/api/courses/{course_id}/knowledge")
        assert framework_response.status_code == 200
        framework = framework_response.json()
        assert framework["version_number"] == 1

        # 顶层是章节容器（旧契约字段保留；文档文件名不再作为容器标题）
        top_titles = [node["title"] for node in framework["tree"]]
        assert top_titles == ["第1章 能量与细胞", "第2章 生态系统"]
        for node in framework["tree"]:
            assert node["is_leaf"] is False
            assert node["sources"] == []
            assert node["children"]

        # 章节下是小节容器（level 2），叶子在 level 3，全部带来源
        chapter = framework["tree"][0]
        assert [child["title"] for child in chapter["children"]] == [
            "1.1 光合作用",
            "1.2 细胞呼吸",
        ]
        leaves = _collect_leaves(framework["tree"])
        assert leaves
        for leaf in leaves:
            assert leaf["is_leaf"] is True
            assert leaf["sources"]
            assert leaf["level"] >= 2

        leaf_node = leaves[0]
        detail_response = await client.get(f"/api/knowledge/{leaf_node['knowledge_node_id']}")
        assert detail_response.status_code == 200
        detail = detail_response.json()
        assert detail["title"]
        assert detail["sources"]
        assert detail["framework_version_id"] == build_body["framework_version_id"]

    reset_settings_cache()


@pytest.mark.anyio
async def test_knowledge_build_one_window_per_chapter(
    migrated_database_url: str,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    pdf_path = tmp_path / "large_lesson.pdf"
    page_texts = []
    for index in range(1, 31):
        if index == 1:
            page_texts.append("第1章 数据结构基础")
        elif index == 11:
            page_texts.append("第2章 树与图")
        elif index == 21:
            page_texts.append("第3章 排序算法")
        else:
            page_texts.append(f"正文内容第 {index} 页，讲解对应章节的概念定义与典型应用。")
    _create_pdf(pdf_path, page_texts)

    section_payloads: list[dict] = []
    outline_calls: list[dict] = []

    class CountingAdapter:
        def ensure_available(self) -> None:
            return None

        def generate_structured_output(self, **kwargs) -> LLMStructuredResult:
            payload = kwargs["user_payload"]
            response_model = kwargs["response_model"]
            if response_model is KnowledgeOutlineV5:
                outline_calls.append(payload)
                data = KnowledgeOutlineV5(
                    root_title="Data Structures",
                    sections=[
                        KnowledgeOutlineSectionV5(
                            ref=f"c{index}",
                            title=f"第{index}章 主题{index}",
                            level=1,
                            parent_ref=None,
                            page_start=start,
                            page_end=start + 9,
                            evidence_page_numbers=[start],
                        )
                        for index, start in ((1, 1), (2, 11), (3, 21))
                    ],
                )
            else:
                section_payloads.append(payload)
                leaves = []
                for segment_index, segment in enumerate(payload["source_segments"][:2], start=1):
                    text = segment["text"]
                    leaves.append(
                        {
                            "ref": f"leaf_{segment_index}",
                            "title": text[:50] or f"Topic {segment_index}",
                            "content": (text[:200] if len(text) >= 24 else text + " padded explanation for the mock extraction."),
                            "sources": [
                                {
                                    "document_id": segment["document_id"],
                                    "page_number": segment["page_start"],
                                    "chunk_id": segment["chunk_id"],
                                    "evidence_type": "chunk" if segment["chunk_id"] else "page_excerpt",
                                }
                            ],
                        }
                    )
                data = KnowledgeSectionExtractionV5(
                    section_ref=payload.get("section_ref", "section"),
                    leaves=leaves,
                )
            return LLMStructuredResult(
                summary=LLMRunSummary(
                    model_name="test-model",
                    prompt_version=kwargs["prompt_version"],
                    provider="test",
                ),
                data=data,
            )

    adapter = CountingAdapter()
    monkeypatch.setattr("app.services.knowledge_service.get_llm_adapter", lambda: adapter)

    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://testserver") as client:
        course_id = await _create_course(client, name="Computer Systems")
        await _upload_pdf_material(client, course_id, pdf_path)

        build_response = await client.post(f"/api/courses/{course_id}/knowledge/build", json={})
        assert build_response.status_code == 202
        build_job = await _wait_for_job(client, build_response.json()["job_id"])
        assert build_job["status"] == "succeeded"

        framework_response = await client.get(f"/api/courses/{course_id}/knowledge")
        framework = framework_response.json()
        leaves = _collect_leaves(framework["tree"])

    # 编号标题被确定性识别，大纲无需 LLM；每个章节一个窗口一次调用
    assert outline_calls == []
    assert len(section_payloads) == 3
    for payload in section_payloads:
        assert payload["max_leaves"] == 10
        assert payload["source_segments"]
        assert all("text" in segment and "document_id" in segment for segment in payload["source_segments"])
    assert [payload["section_path"][0] for payload in section_payloads] == [
        "第1章 数据结构基础",
        "第2章 树与图",
        "第3章 排序算法",
    ]
    assert len(leaves) >= 3


@pytest.mark.anyio
async def test_knowledge_rebuild_uses_section_cache_and_stable_identity(
    migrated_database_url: str,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    pdf_path = tmp_path / "lesson.pdf"
    _create_pdf_with_toc(
        pdf_path,
        [
            "第1章 运输层",
            "用户数据报协议 UDP 提供无连接的不可靠传输服务。",
            "传输控制协议 TCP 提供面向连接的可靠传输服务。",
        ],
        [
            [1, "第1章 运输层", 1],
            [2, "1.1 UDP", 2],
            [2, "1.2 TCP", 3],
        ],
    )

    calls: list[str] = []

    class CountingAdapter:
        def ensure_available(self) -> None:
            return None

        def generate_structured_output(self, **kwargs) -> LLMStructuredResult:
            calls.append(kwargs["prompt_version"])
            payload = kwargs["user_payload"]
            leaves = []
            for segment_index, segment in enumerate(payload["source_segments"][:2], start=1):
                text = segment["text"]
                title = text[:24].strip() or f"Topic {segment_index}"
                leaves.append(
                    {
                        "ref": f"leaf_{segment_index}",
                        "title": title,
                        "content": text[:200] if len(text) >= 24 else text + " padded explanation content.",
                        "sources": [
                            {
                                "document_id": segment["document_id"],
                                "page_number": segment["page_start"],
                                "chunk_id": segment["chunk_id"],
                                "evidence_type": "chunk" if segment["chunk_id"] else "page_excerpt",
                            }
                        ],
                    }
                )
            return LLMStructuredResult(
                summary=LLMRunSummary(
                    model_name="test-model",
                    prompt_version=kwargs["prompt_version"],
                    provider="test",
                ),
                data=KnowledgeSectionExtractionV5(
                    section_ref="section",
                    leaves=leaves,
                ),
            )

    adapter = CountingAdapter()
    monkeypatch.setattr("app.services.knowledge_service.get_llm_adapter", lambda: adapter)

    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://testserver") as client:
        course_id = await _create_course(client, name="Network")
        await _upload_pdf_material(client, course_id, pdf_path)

        first_build = await client.post(f"/api/courses/{course_id}/knowledge/build", json={})
        assert first_build.status_code == 202
        first_job = await _wait_for_job(client, first_build.json()["job_id"])
        assert first_job["status"] == "succeeded"
        first_framework = (await client.get(f"/api/courses/{course_id}/knowledge")).json()
        first_leaf_ids = {leaf["knowledge_node_id"] for leaf in _collect_leaves(first_framework["tree"])}

        calls.clear()
        second_build = await client.post(f"/api/courses/{course_id}/knowledge/build", json={})
        assert second_build.status_code == 202
        second_job = await _wait_for_job(client, second_build.json()["job_id"])
        assert second_job["status"] == "succeeded"
        second_framework = (await client.get(f"/api/courses/{course_id}/knowledge")).json()
        second_leaf_ids = {leaf["knowledge_node_id"] for leaf in _collect_leaves(second_framework["tree"])}

    # 断点缓存：相同内容+prompt+模型的重建不再发起 section 抽取调用
    assert calls == []
    assert second_framework["version_number"] == 2
    assert first_leaf_ids
    assert second_leaf_ids == first_leaf_ids  # 稳定身份：叶子 knowledge_node_id 不变


@pytest.mark.anyio
async def test_duplicate_build_returns_conflict(
    migrated_database_url: str,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setenv("LLM_BASE_URL", "mock://local")
    monkeypatch.setenv("LLM_MODEL", "mock-stage8")
    reset_settings_cache()

    pdf_path = tmp_path / "lesson.pdf"
    _create_pdf(pdf_path, ["Newton's laws describe motion."])

    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://testserver") as client:
        course_id = await _create_course(client, name="Mechanics")
        await _upload_pdf_material(client, course_id, pdf_path)

        version_id = ""
        with session_scope() as session:
            repository = KnowledgeRepository()
            version = repository.create_framework_version(
                session,
                course_id=course_id,
                version_number=1,
                status="building",
            )
            version_id = version.id
            JobRepository().create_job(
                session,
                job_type="knowledge_framework.build",
                payload_json={
                    "course_id": course_id,
                    "framework_version_id": version_id,
                    "reference_outline": None,
                    "job_id": None,
                },
                max_attempts=1,
            )
            session.commit()

        response = await client.post(f"/api/courses/{course_id}/knowledge/build", json={})
        assert response.status_code == 409
        body = response.json()
        assert body["error"]["code"] == "KNOWLEDGE_BUILD_IN_PROGRESS"
        assert body["error"]["details"]["framework_version_id"] == version_id

    reset_settings_cache()


@pytest.mark.anyio
async def test_stale_building_version_self_heals(
    migrated_database_url: str,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setenv("LLM_BASE_URL", "mock://local")
    monkeypatch.setenv("LLM_MODEL", "mock-stage8")
    reset_settings_cache()

    pdf_path = tmp_path / "lesson.pdf"
    _create_pdf_with_toc(
        pdf_path,
        ["第1章 力学基础", "牛顿运动定律描述物体的运动规律与受力关系。"],
        [[1, "第1章 力学基础", 1]],
    )

    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://testserver") as client:
        course_id = await _create_course(client, name="Mechanics")
        await _upload_pdf_material(client, course_id, pdf_path)

        # 模拟进程崩溃遗留的 building 版本（无对应排队任务）
        with session_scope() as session:
            KnowledgeRepository().create_framework_version(
                session,
                course_id=course_id,
                version_number=1,
                status="building",
            )
            session.commit()

        build_response = await client.post(f"/api/courses/{course_id}/knowledge/build", json={})
        assert build_response.status_code == 202
        build_job = await _wait_for_job(client, build_response.json()["job_id"])
        assert build_job["status"] == "succeeded"

        versions_response = await client.get(f"/api/courses/{course_id}/knowledge/versions")
        statuses = [version["status"] for version in versions_response.json()]
        assert statuses == ["active", "failed"]  # 新版本 active，遗留版本被自愈为 failed

    reset_settings_cache()


@pytest.mark.anyio
async def test_knowledge_build_fails_for_invalid_batch_output(
    migrated_database_url: str,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    pdf_path = tmp_path / "lesson.pdf"
    _create_pdf(pdf_path, ["Cache blocks store recently used data."])

    class InvalidStructuredOutputAdapter:
        def ensure_available(self) -> None:
            return None

        def generate_structured_output(self, **_) -> None:
            raise AppError(
                code="LLM_OUTPUT_INVALID",
                message="LLM returned invalid structured output.",
                status_code=502,
            )

    monkeypatch.setattr(
        "app.services.knowledge_service.get_llm_adapter",
        lambda: InvalidStructuredOutputAdapter(),
    )

    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://testserver") as client:
        course_id = await _create_course(client, name="Architecture")
        await _upload_pdf_material(client, course_id, pdf_path)

        build_response = await client.post(f"/api/courses/{course_id}/knowledge/build", json={})
        assert build_response.status_code == 202
        build_job = await _wait_for_job(client, build_response.json()["job_id"])
        assert build_job["status"] == "failed"
        assert build_job["error_code"] == "LLM_OUTPUT_INVALID"
        # D5 修复：构建任务不再自动重跑
        assert build_job["attempts"] == 1


@pytest.mark.anyio
async def test_material_delete_rejected_when_used_by_framework(
    migrated_database_url: str,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setenv("LLM_BASE_URL", "mock://local")
    monkeypatch.setenv("LLM_MODEL", "mock-stage8")
    reset_settings_cache()

    pdf_path = tmp_path / "lesson.pdf"
    _create_pdf_with_toc(
        pdf_path,
        ["第1章 运动定律", "牛顿第一定律说明物体在不受外力时保持静止或匀速直线运动。"],
        [[1, "第1章 运动定律", 1]],
    )

    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://testserver") as client:
        course_id = await _create_course(client, name="Mechanics")
        with pdf_path.open("rb") as handle:
            upload_response = await client.post(
                f"/api/courses/{course_id}/materials",
                files={"file": (pdf_path.name, handle, "application/pdf")},
            )
        material_body = upload_response.json()
        await _wait_for_job(client, material_body["job_id"])

        build_response = await client.post(
            f"/api/courses/{course_id}/knowledge/build",
            json={},
        )
        await _wait_for_job(client, build_response.json()["job_id"])

        delete_response = await client.delete(
            f"/api/materials/{material_body['document_id']}"
        )
        assert delete_response.status_code == 409
        assert delete_response.json()["error"]["code"] == "MATERIAL_IN_USE"

    reset_settings_cache()


@pytest.mark.anyio
async def test_failed_knowledge_rebuild_does_not_replace_active_version(
    migrated_database_url: str,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setenv("LLM_BASE_URL", "mock://local")
    monkeypatch.setenv("LLM_MODEL", "mock-stage8")
    reset_settings_cache()

    pdf_path = tmp_path / "lesson.pdf"
    _create_pdf_with_toc(
        pdf_path,
        ["第1章 生态能量", "能量在生态系统中沿食物链单向流动并逐级递减。"],
        [[1, "第1章 生态能量", 1]],
    )

    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://testserver") as client:
        course_id = await _create_course(client, name="Ecology")
        await _upload_pdf_material(client, course_id, pdf_path)

        first_build_response = await client.post(
            f"/api/courses/{course_id}/knowledge/build",
            json={},
        )
        assert first_build_response.status_code == 202
        first_build_body = first_build_response.json()
        first_job = await _wait_for_job(client, first_build_body["job_id"])
        assert first_job["status"] == "succeeded"

        class FailingAdapter:
            def ensure_available(self) -> None:
                return None

            def generate_structured_output(self, **_) -> None:
                raise AppError(
                    code="LLM_OUTPUT_INVALID",
                    message="Invalid framework output.",
                    status_code=502,
                )

        monkeypatch.setattr(
            "app.services.knowledge_service.get_llm_adapter",
            lambda: FailingAdapter(),
        )

        second_build_response = await client.post(
            f"/api/courses/{course_id}/knowledge/build",
            json={},
        )
        assert second_build_response.status_code == 202
        second_build_body = second_build_response.json()
        second_job = await _wait_for_job(client, second_build_response.json()["job_id"])
        assert second_job["status"] == "failed"
        assert second_job["error_code"] == "LLM_OUTPUT_INVALID"

        versions_response = await client.get(f"/api/courses/{course_id}/knowledge/versions")
        assert versions_response.status_code == 200
        versions = versions_response.json()
        assert [version["status"] for version in versions] == ["failed", "active"]

        framework_response = await client.get(f"/api/courses/{course_id}/knowledge")
        assert framework_response.status_code == 200
        assert framework_response.json()["framework_version_id"] == first_build_body["framework_version_id"]

    reset_settings_cache()


@pytest.mark.anyio
async def test_pptx_build_uses_slide_titles(
    migrated_database_url: str,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setenv("LLM_BASE_URL", "mock://local")
    monkeypatch.setenv("LLM_MODEL", "mock-stage8")
    reset_settings_cache()

    pptx_path = tmp_path / "deck.pptx"
    _create_pptx(
        pptx_path,
        [
            "第1章 网络概述",
            "1.1 网络组成",
            "1.2 网络分类",
        ],
    )

    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://testserver") as client:
        course_id = await _create_course(client, name="Networks")
        await _upload_material(client, course_id, pptx_path, PPTX_MIME)

        build_response = await client.post(f"/api/courses/{course_id}/knowledge/build", json={})
        assert build_response.status_code == 202
        build_job = await _wait_for_job(client, build_response.json()["job_id"])
        assert build_job["status"] == "succeeded"

        framework_response = await client.get(f"/api/courses/{course_id}/knowledge")
        assert framework_response.status_code == 200
        framework = framework_response.json()

        top_titles = [node["title"] for node in framework["tree"]]
        assert top_titles == ["第1章 网络概述"]
        chapter = framework["tree"][0]
        assert [child["title"] for child in chapter["children"]] == [
            "1.1 网络组成",
            "1.2 网络分类",
        ]
        leaves = _collect_leaves(framework["tree"])
        assert leaves
        assert all(leaf["sources"] for leaf in leaves)

    reset_settings_cache()


# --------------------------------------------------------------------------
# v5 修复回归：逐页书签不得变成逐页章节
# --------------------------------------------------------------------------


@pytest.mark.anyio
async def test_knowledge_build_rejects_per_page_bookmarks(
    migrated_database_url: str,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setenv("LLM_BASE_URL", "mock://local")
    monkeypatch.setenv("LLM_MODEL", "mock-stage8")
    reset_settings_cache()

    pdf_path = tmp_path / "slides.pdf"
    page_texts = [f"幻灯片 {index} 的标题与内容" for index in range(1, 31)]
    _create_pdf_with_toc(
        pdf_path,
        page_texts,
        [[1, f"幻灯片 {index} 标题", index] for index in range(1, 31)],
    )

    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://testserver") as client:
        course_id = await _create_course(client, name="Slide Deck")
        await _upload_pdf_material(client, course_id, pdf_path)

        build_response = await client.post(
            f"/api/courses/{course_id}/knowledge/build", json={}
        )
        assert build_response.status_code == 202
        build_job = await _wait_for_job(client, build_response.json()["job_id"])
        assert build_job["status"] == "succeeded"

        framework = (await client.get(f"/api/courses/{course_id}/knowledge")).json()

        # 同一课程重建：大纲与 section 都应命中缓存，不再产生任何 LLM 调用
        rebuild_response = await client.post(
            f"/api/courses/{course_id}/knowledge/build", json={}
        )
        assert rebuild_response.status_code == 202
        rebuild_job = await _wait_for_job(client, rebuild_response.json()["job_id"])
        assert rebuild_job["status"] == "succeeded"

    # 30 条逐页书签被质量闸门拒绝，改由 LLM 兜底给出单一章节，而不是 30 个伪容器
    assert len(framework["tree"]) == 1
    chapter = framework["tree"][0]
    assert chapter["is_leaf"] is False
    assert chapter["children"]
    assert all(node["level"] >= 2 for node in _collect_leaves(framework["tree"]))

    with session_scope() as session:
        rebuild_calls = list(
            session.scalars(select(LLMCall).where(LLMCall.job_id == rebuild_job["id"]))
        )
    assert rebuild_calls == []

    reset_settings_cache()


# --------------------------------------------------------------------------
# v5 修复回归：失败 job 的进度不得把所有单元都算作完成
# --------------------------------------------------------------------------


@pytest.mark.anyio
async def test_failed_build_reports_only_successful_units(
    migrated_database_url: str,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    pdf_path = tmp_path / "lesson.pdf"
    page_texts = []
    for index in range(1, 31):
        if index == 1:
            page_texts.append("第1章 数据结构基础")
        elif index == 11:
            page_texts.append("第2章 树与图")
        elif index == 21:
            page_texts.append("第3章 排序算法")
        else:
            page_texts.append(f"正文内容第 {index} 页，讲解对应章节的概念定义与典型应用。")
    _create_pdf(pdf_path, page_texts)

    class PartialAdapter:
        def ensure_available(self) -> None:
            return None

        def generate_structured_output(self, **kwargs) -> LLMStructuredResult:
            payload = kwargs["user_payload"]
            section_path = payload.get("section_path") or []
            if section_path and section_path[0] == "第1章 数据结构基础":
                segment = payload["source_segments"][0]
                return LLMStructuredResult(
                    summary=LLMRunSummary(
                        model_name="test-model",
                        prompt_version=kwargs["prompt_version"],
                        provider="test",
                    ),
                    data=KnowledgeSectionExtractionV5(
                        section_ref="section",
                        leaves=[
                            {
                                "ref": "leaf_1",
                                "title": "数据结构的基本概念",
                                "content": "数据结构研究数据的逻辑结构、存储结构与基本操作。",
                                "sources": [
                                    {
                                        "document_id": segment["document_id"],
                                        "page_number": segment["page_start"],
                                        "chunk_id": segment["chunk_id"],
                                        "evidence_type": "page_excerpt",
                                    }
                                ],
                            }
                        ],
                    ),
                )
            raise AppError(
                code="LLM_TIMEOUT",
                message="LLM request timed out.",
                status_code=502,
            )

    monkeypatch.setattr(
        "app.services.knowledge_service.get_llm_adapter", lambda: PartialAdapter()
    )

    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://testserver") as client:
        course_id = await _create_course(client, name="Data Structures")
        await _upload_pdf_material(client, course_id, pdf_path)

        build_response = await client.post(
            f"/api/courses/{course_id}/knowledge/build", json={}
        )
        assert build_response.status_code == 202
        build_job = await _wait_for_job(client, build_response.json()["job_id"])

    assert build_job["status"] == "failed"
    assert build_job["error_code"] == "LLM_TIMEOUT"
    # 3 个 section 只有 1 个成功，不能显示 3/3 已完成，也不能到 100%
    assert build_job["total_units"] == 3
    assert build_job["completed_units"] == 1
    assert build_job["progress"] < 70

    with session_scope() as session:
        statuses = list(session.scalars(select(LLMCall.status)))
    assert statuses.count("succeeded") == 1
    assert statuses.count("failed") == 2


# --------------------------------------------------------------------------
# v5 修复回归：预算不足时 job 以 JOB_DEADLINE_EXCEEDED 收尾且不发请求
# --------------------------------------------------------------------------


@pytest.mark.anyio
async def test_knowledge_build_reports_deadline_exceeded(
    migrated_database_url: str,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr("app.services.knowledge_pipeline.BUILD_DEADLINE_SECONDS", 45)
    monkeypatch.setenv("LLM_BASE_URL", "mock://local")
    monkeypatch.setenv("LLM_MODEL", "mock-stage8")
    reset_settings_cache()

    pdf_path = tmp_path / "lesson.pdf"
    page_texts = []
    for index in range(1, 31):
        if index == 1:
            page_texts.append("第1章 数据结构基础")
        elif index == 11:
            page_texts.append("第2章 树与图")
        elif index == 21:
            page_texts.append("第3章 排序算法")
        else:
            page_texts.append(f"正文内容第 {index} 页，讲解对应章节的概念定义与典型应用。")
    _create_pdf(pdf_path, page_texts)

    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://testserver") as client:
        course_id = await _create_course(client, name="Deadline")
        await _upload_pdf_material(client, course_id, pdf_path)

        build_response = await client.post(
            f"/api/courses/{course_id}/knowledge/build", json={}
        )
        assert build_response.status_code == 202
        build_job = await _wait_for_job(client, build_response.json()["job_id"])

    assert build_job["status"] == "failed"
    assert build_job["error_code"] == "JOB_DEADLINE_EXCEEDED"
    assert build_job["completed_units"] == 0
    assert build_job["progress"] < 70

    with session_scope() as session:
        rows = list(session.scalars(select(LLMCall)))
    assert [row.status for row in rows] == ["skipped_deadline"]
    assert rows[0].attempt == 0
    assert rows[0].latency_ms is None

    reset_settings_cache()


# --------------------------------------------------------------------------
# 文档级伞形章：单章文档被拆成一堆"一级节"时，合并后顶层 = 文档级章数
# --------------------------------------------------------------------------


@pytest.mark.anyio
async def test_single_chapter_document_is_wrapped_into_umbrella_chapter(
    migrated_database_url: str,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    pdf_path = tmp_path / "ch1_引言_v2.0.pdf"
    page_texts = []
    for index in range(1, 21):
        page_texts.append(
            "第一章 引言"
            if index == 1
            else f"正文内容第 {index} 页，讲解计算机网络的基本概念与典型应用。"
        )
    _create_pdf(pdf_path, page_texts)

    outline_calls: list[dict] = []
    section_payloads: list[dict] = []

    class SectionLevelOutlineAdapter:
        """模拟 Stage A 把文档内小节当成一级章：两个 L1、没有 L2。"""

        def ensure_available(self) -> None:
            return None

        def generate_structured_output(self, **kwargs) -> LLMStructuredResult:
            payload = kwargs["user_payload"]
            if kwargs["response_model"] is KnowledgeOutlineV5:
                outline_calls.append(payload)
                data = KnowledgeOutlineV5(
                    root_title="第一章 引言",
                    sections=[
                        KnowledgeOutlineSectionV5(
                            ref="s1",
                            title="1.1 背景",
                            level=1,
                            parent_ref=None,
                            page_start=1,
                            page_end=10,
                            evidence_page_numbers=[1],
                        ),
                        KnowledgeOutlineSectionV5(
                            ref="s2",
                            title="1.2 方法",
                            level=1,
                            parent_ref=None,
                            page_start=11,
                            page_end=20,
                            evidence_page_numbers=[11],
                        ),
                    ],
                )
            else:
                section_payloads.append(payload)
                segment = payload["source_segments"][0]
                data = KnowledgeSectionExtractionV5(
                    section_ref=payload.get("section_ref", "section"),
                    leaves=[
                        {
                            "ref": "leaf_1",
                            "title": f"知识点 {len(section_payloads)}",
                            "content": "这一节的核心概念与典型应用。" * 3,
                            "sources": [
                                {
                                    "document_id": segment["document_id"],
                                    "page_number": segment["page_start"],
                                    "chunk_id": segment["chunk_id"],
                                    "evidence_type": "chunk" if segment["chunk_id"] else "page_excerpt",
                                }
                            ],
                        }
                    ],
                )
            return LLMStructuredResult(
                summary=LLMRunSummary(
                    model_name="test-model",
                    prompt_version=kwargs["prompt_version"],
                    provider="test",
                ),
                data=data,
            )

    adapter = SectionLevelOutlineAdapter()
    monkeypatch.setattr("app.services.knowledge_service.get_llm_adapter", lambda: adapter)

    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://testserver") as client:
        course_id = await _create_course(client, name="Umbrella")
        await _upload_pdf_material(client, course_id, pdf_path)

        build_response = await client.post(f"/api/courses/{course_id}/knowledge/build", json={})
        assert build_response.status_code == 202
        build_job = await _wait_for_job(client, build_response.json()["job_id"])
        assert build_job["status"] == "succeeded"

        framework = (await client.get(f"/api/courses/{course_id}/knowledge")).json()

    # 两个小节级 L1 被收进 root_title 指定的伞形章，顶层只有 1 个
    assert len(outline_calls) == 1
    assert len(framework["tree"]) == 1
    umbrella = framework["tree"][0]
    assert umbrella["title"] == "第一章 引言"
    assert [child["title"] for child in umbrella["children"]] == ["1.1 背景", "1.2 方法"]
    assert [payload["section_path"][0] for payload in section_payloads] == [
        "第一章 引言",
        "第一章 引言",
    ]
    # 首节吸收章首：1.1 原本从第 1 页起，伞形章起点也是 1，窗口应覆盖前 10 页
    assert section_payloads[0]["source_segments"][0]["page_start"] == 1
    assert len(_collect_leaves(framework["tree"])) >= 2

    reset_settings_cache()


# --------------------------------------------------------------------------
# Phase 2：类型化关系（落库 → 详情接口 → 图端点）
# --------------------------------------------------------------------------


@pytest.mark.anyio
async def test_relations_are_persisted_and_exposed(
    migrated_database_url: str,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setenv("LLM_BASE_URL", "mock://local")
    monkeypatch.setenv("LLM_MODEL", "mock-stage8")
    reset_settings_cache()

    pdf_path = tmp_path / "transport.pdf"
    _create_pdf_with_toc(
        pdf_path,
        [
            "第1章 运输层",
            "用户数据报协议 UDP 提供无连接的不可靠传输服务。",
            "传输控制协议 TCP 提供面向连接的可靠传输服务。",
        ],
        [[1, "第1章 运输层", 1]],
    )

    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://testserver") as client:
        course_id = await _create_course(client, name="Transport")
        await _upload_pdf_material(client, course_id, pdf_path)

        build_response = await client.post(
            f"/api/courses/{course_id}/knowledge/build", json={}
        )
        assert build_response.status_code == 202
        build_body = build_response.json()
        build_job = await _wait_for_job(client, build_body["job_id"])
        assert build_job["status"] == "succeeded"

        framework = (await client.get(f"/api/courses/{course_id}/knowledge")).json()
        leaves = _collect_leaves(framework["tree"])
        assert len(leaves) >= 2  # mock 每个 page 一个叶子，才能构成关系
        leaf_ids = {leaf["knowledge_node_id"] for leaf in leaves}

        detail_with_relation = None
        for leaf in leaves:
            detail = (
                await client.get(f"/api/knowledge/{leaf['knowledge_node_id']}")
            ).json()
            if detail["relations"]:
                detail_with_relation = detail
                break
        assert detail_with_relation is not None
        for relation in detail_with_relation["relations"]:
            assert relation["relation_type"] in RELATION_TYPES
            assert 0.0 <= relation["confidence"] <= 1.0
            assert relation["target_knowledge_node_id"] in leaf_ids
            assert relation["target_title"]

        graph_response = await client.get(f"/api/courses/{course_id}/knowledge/graph")
        assert graph_response.status_code == 200
        graph = graph_response.json()
        assert graph["framework_version_id"] == build_body["framework_version_id"]
        assert len(graph["edges"]) == 2
        node_ids = {node["knowledge_node_id"] for node in graph["nodes"]}
        assert node_ids <= leaf_ids
        for edge in graph["edges"]:
            assert edge["source"] in node_ids
            assert edge["target"] in node_ids
            assert edge["source"] != edge["target"]
            assert edge["relation_type"] in RELATION_TYPES
        assert all(node["section_path"] == ["第1章 运输层"] for node in graph["nodes"])

        # 关系随 section 缓存一起复用：重建 0 次调用，边仍然完整
        rebuild_response = await client.post(
            f"/api/courses/{course_id}/knowledge/build", json={}
        )
        assert rebuild_response.status_code == 202
        rebuild_job = await _wait_for_job(client, rebuild_response.json()["job_id"])
        assert rebuild_job["status"] == "succeeded"
        rebuilt_graph = (
            await client.get(f"/api/courses/{course_id}/knowledge/graph")
        ).json()
        assert _graph_signature(rebuilt_graph) == _graph_signature(graph)

    with session_scope() as session:
        rows = list(session.scalars(select(KnowledgeNodeRelation)))
    assert len(rows) == 4  # 两个版本各 2 条：关系是版本作用域、随框架重写
    assert {row.relation_type for row in rows} <= set(RELATION_TYPES)

    reset_settings_cache()


def _graph_signature(graph: dict) -> tuple[set[tuple[str, str, str]], set[str]]:
    edges = {
        (edge["source"], edge["target"], edge["relation_type"])
        for edge in graph["edges"]
    }
    nodes = {node["knowledge_node_id"] for node in graph["nodes"]}
    return edges, nodes


@pytest.mark.anyio
async def test_invalid_relations_do_not_fail_the_build(
    migrated_database_url: str,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    pdf_path = tmp_path / "lesson.pdf"
    _create_pdf_with_toc(
        pdf_path,
        ["第1章 基础", "缓存块保存最近使用的数据。", "缺页中断由操作系统处理。"],
        [[1, "第1章 基础", 1]],
    )

    class GarbageRelationsAdapter:
        """叶子合法、关系全部非法：构建必须成功、关系全部丢弃。"""

        def ensure_available(self) -> None:
            return None

        def generate_structured_output(self, **kwargs) -> LLMStructuredResult:
            payload = kwargs["user_payload"]
            segment = payload["source_segments"][0]
            data = KnowledgeSectionExtractionV5.model_validate(
                {
                    "section_ref": "section",
                    "leaves": [
                        {
                            "ref": "leaf_1",
                            "title": "缓存块",
                            "content": "缓存块保存最近使用的数据，减少重复访问慢速存储。",
                            "sources": [
                                {
                                    "document_id": segment["document_id"],
                                    "page_number": segment["page_start"],
                                    "chunk_id": segment["chunk_id"],
                                    "evidence_type": "page_excerpt",
                                }
                            ],
                        },
                        {
                            "ref": "leaf_2",
                            "title": "缺页中断",
                            "content": "缺页中断由操作系统处理，把缺失的页面从磁盘调入内存。",
                            "sources": [
                                {
                                    "document_id": segment["document_id"],
                                    "page_number": segment["page_start"],
                                    "chunk_id": segment["chunk_id"],
                                    "evidence_type": "page_excerpt",
                                }
                            ],
                        },
                    ],
                    "relations": [
                        # 引用不存在的叶子 + 置信度非法
                        {
                            "source_ref": "leaf_1",
                            "target_ref": "leaf_9",
                            "relation_type": "prerequisite",
                            "confidence": "high",
                            "evidence": {"text": "not a string"},
                        },
                        # 自环
                        {
                            "source_ref": "leaf_1",
                            "target_ref": "leaf_1",
                            "relation_type": "causes",
                            "confidence": 0.5,
                        },
                        # 未知类型
                        {
                            "source_ref": "leaf_1",
                            "target_ref": "leaf_2",
                            "relation_type": "related",
                            "confidence": 0.5,
                        },
                    ],
                }
            )
            return LLMStructuredResult(
                summary=LLMRunSummary(
                    model_name="test-model",
                    prompt_version=kwargs["prompt_version"],
                    provider="test",
                ),
                data=data,
            )

    monkeypatch.setattr(
        "app.services.knowledge_service.get_llm_adapter",
        lambda: GarbageRelationsAdapter(),
    )

    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://testserver") as client:
        course_id = await _create_course(client, name="Memory")
        await _upload_pdf_material(client, course_id, pdf_path)

        build_response = await client.post(
            f"/api/courses/{course_id}/knowledge/build", json={}
        )
        assert build_response.status_code == 202
        build_job = await _wait_for_job(client, build_response.json()["job_id"])
        assert build_job["status"] == "succeeded"

        framework = (await client.get(f"/api/courses/{course_id}/knowledge")).json()
        leaves = _collect_leaves(framework["tree"])
        assert len(leaves) == 2
        detail = (
            await client.get(f"/api/knowledge/{leaves[0]['knowledge_node_id']}")
        ).json()
        assert detail["relations"] == []

        graph = (await client.get(f"/api/courses/{course_id}/knowledge/graph")).json()
        assert graph["nodes"] == []
        assert graph["edges"] == []

    with session_scope() as session:
        assert list(session.scalars(select(KnowledgeNodeRelation))) == []

    reset_settings_cache()
